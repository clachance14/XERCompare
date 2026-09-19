"""Verify the desktop away from source/Python PATH; optional developer CLI check."""

from __future__ import annotations

import hashlib
import argparse
import json
import os
from pathlib import Path
import shutil
import struct
import subprocess
from tempfile import TemporaryDirectory
import tomllib

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "build" / "verification"


def inspect_exe(path: Path, subsystem: int):
    from PyInstaller.archive.readers import CArchiveReader

    data = path.read_bytes()
    offset = struct.unpack_from("<I", data, 0x3C)[0]
    assert data[offset : offset + 4] == b"PE\0\0"
    assert struct.unpack_from("<H", data, offset + 4)[0] == 0x8664
    assert struct.unpack_from("<H", data, offset + 24 + 68)[0] == subsystem
    archive = CArchiveReader(str(path))
    names = list(archive.toc)
    assert "python313.dll" in names
    assert not any("darkgarden" in name.lower() for name in names)
    assert any("THIRD_PARTY_NOTICES.txt" in name for name in names)
    assert any(name.endswith("Vera.ttf") for name in names)
    if subsystem == 2:
        assert any("_tkinter.pyd" in name for name in names)
        assert any(name.endswith("tk.tcl") for name in names)
    embedded = archive.open_embedded_archive("PYZ.pyz")
    assert not any(
        name.split(".")[0] in {"PyP6Xer", "pyp6xer", "numpy", "pandas", "requests"}
        for name in embedded.toc
    )
    return {
        "file": path.name,
        "bytes": len(data),
        "sha256": hashlib.sha256(data).hexdigest(),
        "architecture": "Windows x64",
        "subsystem": "windowed" if subsystem == 2 else "console",
        "bundled_python": "python313.dll",
        "archive_entries": len(names),
    }


def check_reports(folder):
    from openpyxl import load_workbook

    book = load_workbook(folder / "compare.xlsx", read_only=True, data_only=True)
    try:
        assert len(book.sheetnames) == 39
        assert book["Changes Combined"].max_row - 4 == 29
    finally:
        book.close()
    assert (folder / "compare.pdf").read_bytes().startswith(b"%PDF-")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--with-cli", action="store_true")
    args = parser.parse_args()
    if os.name != "nt":
        raise SystemExit("Use Windows Python to verify the Windows executables.")
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    binaries = [inspect_exe(ROOT / "dist" / "XERCompare.exe", 2)]
    if args.with_cli:
        binaries.append(inspect_exe(ROOT / "dist" / "XERCompare-CLI.exe", 3))
    # Keep required Windows user/session variables, but remove Python and dev PATHs.
    env = {
        key: value
        for key, value in os.environ.items()
        if not key.upper().startswith(("PYTHON", "VIRTUAL_ENV", "CONDA", "_PYI"))
    }
    windows = env.get("SystemRoot", r"C:\Windows")
    env["PATH"] = f"{windows}\\System32;{windows}"
    env["PYTHONNOUSERSITE"] = "1"
    with TemporaryDirectory(prefix="XERCompare-isolated-") as directory:
        isolated = Path(directory)
        for name in (binary["file"] for binary in binaries):
            shutil.copy2(ROOT / "dist" / name, isolated / name)
        for name in ("base.xer", "revised.xer"):
            shutil.copy2(ROOT / "demo" / "phase3" / name, isolated / name)
        common = dict(
            cwd=isolated, env=env, timeout=360, capture_output=True, text=True
        )
        gui = subprocess.run(
            [
                str(isolated / "XERCompare.exe"),
                "--self-test",
                "base.xer",
                "revised.xer",
                "-o",
                "gui-output",
            ],
            **common,
        )
        assert gui.returncode == 0, (gui.returncode, gui.stdout, gui.stderr)
        record = json.loads(
            (isolated / "gui-output" / "standalone-verification.json").read_text()
        )
        assert record["ok"] and record["frozen"], record
        assert (
            record["quality_rows"] > 0
            and record["history_pairs"] == 2
            and record["profile_verified"]
        ), record
        assert record["reports"] == 39 and record["preview_rows"] == 29, record
        assert record["summary"]["modified"] == 2, record
        assert record["summary"]["added"] == record["summary"]["deleted"] == 1
        check_reports(Path(record["folder"]))
        extra = {}
        outputs = ["gui-output"]
        if args.with_cli:
            cli = subprocess.run(
                [
                    str(isolated / "XERCompare-CLI.exe"),
                    "base.xer",
                    "revised.xer",
                    "-o",
                    "cli-output",
                ],
                **common,
            )
            assert cli.returncode == 0, (cli.returncode, cli.stdout, cli.stderr)
            check_reports(isolated / "cli-output")
            version = subprocess.run(
                [str(isolated / "XERCompare-CLI.exe"), "--version"], **common
            )
            expected_version = tomllib.loads((ROOT / "pyproject.toml").read_text())[
                "project"
            ]["version"]
            assert version.returncode == 0 and expected_version in version.stdout
            (isolated / "bad.xer").write_text("Not an XER", encoding="cp1252")
            failure = subprocess.run(
                [
                    str(isolated / "XERCompare-CLI.exe"),
                    "bad.xer",
                    "revised.xer",
                    "-o",
                    "bad-output",
                ],
                **common,
            )
            assert failure.returncode == 2 and "error:" in failure.stderr
            extra = {
                "cli_version": version.stdout.strip(),
                "malformed_input_exit": failure.returncode,
            }
            outputs.append("cli-output")
            (EVIDENCE / "cli-output.txt").write_text(cli.stdout, encoding="utf-8")
        for name in outputs:
            shutil.copytree(isolated / name, EVIDENCE / name, dirs_exist_ok=True)
        (EVIDENCE / "release-verification.json").write_text(
            json.dumps(
                {
                    "ok": True,
                    "binaries": binaries,
                    "isolated_path": env["PATH"],
                    "gui": record,
                    **extra,
                    "limit": "Verified on this Windows host with isolated PATH and bundled module origins; no separate Python-free VM was available.",
                },
                indent=2,
            ),
            encoding="utf-8",
        )
    print(f"Standalone desktop verified. Evidence: {EVIDENCE}")


if __name__ == "__main__":
    main()
