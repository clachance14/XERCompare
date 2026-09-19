from pathlib import Path

from openpyxl import load_workbook
import pytest

from tests.fixtures.make_xer import make_xer
from xercompare.compare.engine import compare_schedules
from xercompare.model.schedule import build_schedule
from xercompare.parse.xer_parser import parse_xer
from xercompare.reports.excel import MOD_FILL, write_excel


def schedule(path, **kwargs):
    path.write_text(make_xer(**kwargs), encoding="cp1252")
    return build_schedule(parse_xer(path))


def workbook(tmp_path, left, right):
    comparison = compare_schedules(
        schedule(tmp_path / "left.xer", **left),
        schedule(tmp_path / "right.xer", **right),
    )
    return comparison, load_workbook(write_excel(comparison, tmp_path / "compare.xlsx"))


def records(ws):
    headers = [c.value for c in ws[4]]
    return [
        dict(zip(headers, row)) for row in ws.iter_rows(min_row=5, values_only=True)
    ]


def assert_grid(ws):
    assert ws.freeze_panes == "A5"
    assert ws.auto_filter.ref == f"A4:{ws.cell(ws.max_row, ws.max_column).coordinate}"


def test_inventory_sheets(tmp_path):
    c, wb = workbook(
        tmp_path,
        {
            "activities": [
                {"task_code": "DELETED", "task_name": "Old"},
                {"task_code": "UNCHANGED", "task_name": "Same"},
            ]
        },
        {
            "activities": [
                {"task_id": "20", "task_code": "UNCHANGED", "task_name": "Same"},
                {"task_code": "ADDED", "task_name": "New"},
            ]
        },
    )
    assert c.summary == dict(
        left_activities=2,
        right_activities=2,
        added=1,
        deleted=1,
        modified=0,
        unchanged=1,
        logic_added=0,
        logic_deleted=0,
        logic_modified=0,
    )
    assert wb["Summary"]["B7"].value == "DEMO"
    for title, code in [("Added Tasks", "ADDED"), ("Deleted Tasks", "DELETED")]:
        ws = wb[title]
        assert [c.value for c in ws[4]][:3] == ["Activity ID", "Name", "WBS"]
        assert [r["Activity ID"] for r in records(ws)] == [code]
        assert records(ws)[0]["WBS"] == "DEMO"
        assert_grid(ws)


def check_field_sheet(
    tmp_path, title, left_fields, right_fields, expected_pairs, klass
):
    base = {
        "task_id": "1",
        "task_code": "CHANGED",
        "task_name": "Work",
        "remain_hr": 40,
        **left_fields,
    }
    revised = {**base, "task_id": "91", **right_fields}
    c, wb = workbook(
        tmp_path,
        {
            "activities": [
                base,
                {"task_code": "UNCHANGED", "task_name": "Same"},
                {"task_code": "OTHER", "task_name": "Old name"},
                {"task_code": "DELETED"},
            ],
            "calendars": CALENDARS,
        },
        {
            "activities": [
                revised,
                {"task_code": "UNCHANGED", "task_name": "Same"},
                {"task_code": "OTHER", "task_name": "New name"},
                {"task_code": "ADDED"},
            ],
            "calendars": CALENDARS,
        },
    )
    ws = wb[title]
    rows = records(ws)
    # The description sheet also intentionally captures OTHER's rename.
    expected_ids = (
        ["CHANGED", "OTHER"] if title == "Duration Description Change" else ["CHANGED"]
    )
    assert [r["Activity ID"] for r in rows] == expected_ids
    assert rows[0]["WBS"] == "DEMO"
    assert next(d for d in c.modified if d.task_code == "CHANGED").klass == klass
    headers = [cell.value for cell in ws[4]]
    for label, (before, after) in expected_pairs.items():
        assert rows[0][f"Base {label}"] == before
        assert rows[0][f"Revised {label}"] == after
    for col, header in enumerate(headers, 1):
        if header.startswith("Revised "):
            changed = rows[0][header] != rows[0][header.replace("Revised ", "Base ", 1)]
            assert (ws.cell(5, col).fill == MOD_FILL) == changed
    assert_grid(ws)


CALENDARS = [
    {"clndr_id": "1", "clndr_name": "Standard 5x8", "day_hr_cnt": 8},
    {"clndr_id": "2", "clndr_name": "Six by ten", "day_hr_cnt": 10},
]


@pytest.mark.parametrize(
    "field,label", [("act_start", "Actual start"), ("act_end", "Actual finish")]
)
def test_actual_dates(tmp_path, field, label):
    check_field_sheet(
        tmp_path,
        "Actual Date Change",
        {field: ""},
        {field: "2024-02-02 08:00"},
        {label: (None, "2024-02-02 08:00")},
        "progress",
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        (
            {"remain_hr": 40},
            {"remain_hr": 16},
            {"Remaining duration (days)": (5, 2)},
            "progress",
        ),
        ({"pct": 0}, {"pct": 25}, {"Physical complete (%)": (0, 25)}, "progress"),
        (
            {"status_code": "TK_NotStart"},
            {"status_code": "TK_Active"},
            {"Status": ("TK_NotStart", "TK_Active")},
            "progress",
        ),
        (
            {"restart_date": ""},
            {"restart_date": "2024-02-02 08:00"},
            {"Restart": (None, "2024-02-02 08:00")},
            "progress",
        ),
        (
            {"reend_date": ""},
            {"reend_date": "2024-02-02 08:00"},
            {"Re-end": (None, "2024-02-02 08:00")},
            "progress",
        ),
        (
            {"suspend_date": ""},
            {"suspend_date": "2024-02-02 08:00"},
            {"Suspend": (None, "2024-02-02 08:00")},
            "progress",
        ),
        (
            {"resume_date": ""},
            {"resume_date": "2024-02-02 08:00"},
            {"Resume": (None, "2024-02-02 08:00")},
            "progress",
        ),
    ],
)
def test_progress_change(tmp_path, left_fields, right_fields, expected_pairs, klass):
    check_field_sheet(
        tmp_path, "Progress Change", left_fields, right_fields, expected_pairs, klass
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        (
            {"dur_hr": 40},
            {"dur_hr": 80},
            {"Original duration (days)": (5, 10)},
            "revision",
        ),
        (
            {"task_name": "Old"},
            {"task_name": "New"},
            {"Description": ("Old", "New")},
            "revision",
        ),
    ],
)
def test_duration_description_change(
    tmp_path, left_fields, right_fields, expected_pairs, klass
):
    check_field_sheet(
        tmp_path,
        "Duration Description Change",
        left_fields,
        right_fields,
        expected_pairs,
        klass,
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        (
            {"clndr_id": "1"},
            {"clndr_id": "2"},
            {"Calendar": ("Standard 5x8", "Six by ten")},
            "both",
        )
    ],
)
def test_calendar_change(tmp_path, left_fields, right_fields, expected_pairs, klass):
    check_field_sheet(
        tmp_path, "Calendar Change", left_fields, right_fields, expected_pairs, klass
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        (
            {"cstr_type": ""},
            {"cstr_type": "CS_MSO"},
            {"Constraint type": (None, "CS_MSO")},
            "revision",
        ),
        (
            {"cstr_date": ""},
            {"cstr_date": "2024-02-02 08:00"},
            {"Constraint date": (None, "2024-02-02 08:00")},
            "revision",
        ),
        (
            {"cstr_type2": ""},
            {"cstr_type2": "CS_MSO"},
            {"Secondary constraint type": (None, "CS_MSO")},
            "revision",
        ),
        (
            {"cstr_date2": ""},
            {"cstr_date2": "2024-02-02 08:00"},
            {"Secondary constraint date": (None, "2024-02-02 08:00")},
            "revision",
        ),
    ],
)
def test_constraint_change(tmp_path, left_fields, right_fields, expected_pairs, klass):
    check_field_sheet(
        tmp_path, "Constraint Change", left_fields, right_fields, expected_pairs, klass
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        ({"tf_hr": 80}, {"tf_hr": -16}, {"Total float (days)": (10, -2)}, "date"),
        ({"ff_hr": ""}, {"ff_hr": 0}, {"Free float (days)": (None, 0)}, "date"),
    ],
)
def test_float_change(tmp_path, left_fields, right_fields, expected_pairs, klass):
    check_field_sheet(
        tmp_path, "Float Change", left_fields, right_fields, expected_pairs, klass
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        (
            {"es": "2024-01-08 08:00"},
            {"es": "2024-02-02 08:00"},
            {"Early start": ("2024-01-08 08:00", "2024-02-02 08:00")},
            "date",
        )
    ],
)
def test_early_start_changes(
    tmp_path, left_fields, right_fields, expected_pairs, klass
):
    check_field_sheet(
        tmp_path,
        "Early Start Changes",
        left_fields,
        right_fields,
        expected_pairs,
        klass,
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        (
            {"ef": "2024-01-08 08:00"},
            {"ef": "2024-02-02 08:00"},
            {"Early finish": ("2024-01-08 08:00", "2024-02-02 08:00")},
            "date",
        )
    ],
)
def test_early_finish_changes(
    tmp_path, left_fields, right_fields, expected_pairs, klass
):
    check_field_sheet(
        tmp_path,
        "Early Finish Changes",
        left_fields,
        right_fields,
        expected_pairs,
        klass,
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        (
            {"ls": "2024-01-08 08:00"},
            {"ls": "2024-02-02 08:00"},
            {"Late start": ("2024-01-08 08:00", "2024-02-02 08:00")},
            "date",
        )
    ],
)
def test_late_start_changes(tmp_path, left_fields, right_fields, expected_pairs, klass):
    check_field_sheet(
        tmp_path, "Late Start Changes", left_fields, right_fields, expected_pairs, klass
    )


@pytest.mark.parametrize(
    "left_fields,right_fields,expected_pairs,klass",
    [
        (
            {"lf": "2024-01-08 08:00"},
            {"lf": "2024-02-02 08:00"},
            {"Late finish": ("2024-01-08 08:00", "2024-02-02 08:00")},
            "date",
        )
    ],
)
def test_late_finish_changes(
    tmp_path, left_fields, right_fields, expected_pairs, klass
):
    check_field_sheet(
        tmp_path,
        "Late Finish Changes",
        left_fields,
        right_fields,
        expected_pairs,
        klass,
    )


def test_variances_are_complete_numeric_and_not_limited_to_500(tmp_path):
    activities = [
        {
            "task_code": f"A{i:04}",
            "task_name": "Work",
            "remain_hr": 40,
            "ls": "2024-01-08 08:00",
            "lf": "2024-01-12 17:00",
            "ff_hr": 16,
        }
        for i in range(505)
    ]
    revised = [{**a, "dur_hr": 80} for a in activities]
    _, wb = workbook(tmp_path, {"activities": activities}, {"activities": revised})
    ws = wb["Variances"]
    rows = records(ws)
    assert len(rows) == 505
    assert rows[-1]["Activity ID"] == "A0504"
    for label, expected in [
        ("Early start", "2024-01-08 08:00"),
        ("Early finish", "2024-01-12 17:00"),
        ("Late start", "2024-01-08 08:00"),
        ("Late finish", "2024-01-12 17:00"),
        ("Total float (days)", 10),
        ("Free float (days)", 2),
        ("Original duration (days)", 5),
    ]:
        assert rows[0][f"Base {label}"] == expected
    assert rows[0]["Revised Original duration (days)"] == 10
    assert_grid(ws)


def logic_pair():
    left = {
        "activities": [
            {"task_id": str(i), "task_code": f"A{i}", "task_name": f"Task {i}"}
            for i in range(1, 4)
        ],
        "preds": [
            {"pred_id": "1", "succ_id": "2", "type": "PR_FS", "lag": 0},
            {"pred_id": "2", "succ_id": "3", "type": "PR_FS", "lag": 0},
        ],
    }
    right = {
        "activities": [
            {"task_id": str(i + 10), "task_code": f"A{i}", "task_name": f"Task {i}"}
            for i in range(1, 4)
        ],
        "preds": [
            {"pred_id": "11", "succ_id": "12", "type": "PR_SS", "lag": 8},
            {"pred_id": "12", "succ_id": "13", "type": "PR_FS", "lag": 16},
        ],
    }
    return left, right


def test_added_relationships(tmp_path):
    left, right = logic_pair()
    c, wb = workbook(tmp_path, left, right)
    ws = wb["Added Relationships"]
    (row,) = records(ws)
    assert (
        row["Activity ID"],
        row["Predecessor"],
        row["Base Lag (hours)"],
        row["Revised Lag (hours)"],
    ) == ("A2", "A1", None, 8)
    assert row["WBS"] == "DEMO"
    assert row["Class"] == "revision"
    assert len(c.logic_added) == 1
    assert not c.modified  # Network edits have their own delta counts.
    assert_grid(ws)


def test_deleted_relationships(tmp_path):
    left, right = logic_pair()
    c, wb = workbook(tmp_path, left, right)
    ws = wb["Deleted Relationships"]
    (row,) = records(ws)
    assert (
        row["Activity ID"],
        row["Predecessor"],
        row["Base Lag (hours)"],
        row["Revised Lag (hours)"],
    ) == ("A2", "A1", 0, None)
    assert row["WBS"] == "DEMO"
    assert row["Class"] == "revision"
    assert len(c.logic_deleted) == 1
    assert not c.modified  # Network edits have their own delta counts.
    assert_grid(ws)


def test_lag_change(tmp_path):
    left, right = logic_pair()
    c, wb = workbook(tmp_path, left, right)
    ws = wb["Lag Change"]
    (row,) = records(ws)
    assert (
        row["Activity ID"],
        row["Predecessor"],
        row["Base Lag (hours)"],
        row["Revised Lag (hours)"],
    ) == ("A3", "A2", 0, 16)
    assert row["WBS"] == "DEMO"
    assert row["Class"] == "revision"
    assert len(c.logic_modified) == 1
    assert not c.modified  # Network edits have their own delta counts.
    assert_grid(ws)


def test_driving_path_is_stored_float_approximation(tmp_path):
    base = [
        {"task_code": code, "tf_hr": value}
        for code, value in [
            ("ENTER", 8),
            ("LEAVE", 0),
            ("STAY", -8),
            ("FROM_UNKNOWN", ""),
            ("TO_UNKNOWN", 0),
            ("NEVER", 8),
            ("DELETED", 0),
        ]
    ]
    revised = [
        {"task_code": code, "tf_hr": value}
        for code, value in [
            ("ENTER", 0),
            ("LEAVE", 8),
            ("STAY", -8),
            ("FROM_UNKNOWN", 0),
            ("TO_UNKNOWN", ""),
            ("NEVER", 16),
            ("ADDED", -8),
        ]
    ]
    _, wb = workbook(tmp_path, {"activities": base}, {"activities": revised})
    ws = wb["Driving Path Change"]
    assert "approximation" in ws["A2"].value.lower()
    assert "stored" in ws["A2"].value.lower()
    rows = {r["Activity ID"]: r for r in records(ws)}
    assert set(rows) == {
        "ENTER",
        "LEAVE",
        "STAY",
        "FROM_UNKNOWN",
        "TO_UNKNOWN",
        "DELETED",
        "ADDED",
    }
    assert rows["ENTER"]["Change type"] == "Entered"
    assert rows["LEAVE"]["Change type"] == "Left"
    assert rows["STAY"]["Change type"] == "Retained"
    assert rows["FROM_UNKNOWN"]["Base Membership"] == "Unknown"
    assert rows["TO_UNKNOWN"]["Revised Membership"] == "Unknown"
    assert rows["ADDED"]["Base Membership"] == "Absent"
    assert rows["DELETED"]["Revised Membership"] == "Absent"
    assert_grid(ws)


def resource_pair():
    left = {
        "activities": [
            {"task_id": "1", "task_code": "A1"},
            {"task_id": "2", "task_code": "A2"},
        ],
        "resources": [
            {"rsrc_id": "1", "rsrc_name": "Crew"},
            {"rsrc_id": "2", "rsrc_name": "Gone"},
        ],
        "assignments": [
            {"task_id": "1", "rsrc_id": "1", "target_qty": 40},
            {"task_id": "1", "rsrc_id": "1", "target_qty": 50},
            {"task_id": "2", "rsrc_id": "2", "target_cost": 100},
        ],
    }
    right = {
        "activities": [
            {"task_id": "11", "task_code": "A1"},
            {"task_id": "12", "task_code": "A2"},
        ],
        "resources": [
            {"rsrc_id": "91", "rsrc_name": "Crew"},
            {"rsrc_id": "92", "rsrc_name": "New"},
        ],
        "assignments": [
            {"task_id": "11", "rsrc_id": "91", "target_qty": 80},
            {"task_id": "12", "rsrc_id": "92", "target_cost": 200},
        ],
    }
    return left, right


def test_added_resource_assignments(tmp_path):
    c, wb = workbook(tmp_path, *resource_pair())
    ws = wb["Added Resource Assignments"]
    (row,) = records(ws)
    assert row["Activity ID"] == "A2"
    assert row["Revised Resource"] == "New"
    assert row["Revised Target cost"] == 200
    assert row["Base Resource"] is None
    assert len(c.resource_added) == 1
    assert not c.modified  # Assignment inventory is separate from TASK fields.
    assert_grid(ws)


def test_resource_assignment_id_fallback_and_project_scope(tmp_path):
    left = {
        "activities": [{"task_id": "1", "task_code": "A1"}],
        "resources": [{"rsrc_id": "1", "rsrc_name": "Old name"}],
        "assignments": [{"proj_id": "1", "task_id": "1", "rsrc_id": "1"}],
    }
    right = {
        "activities": [
            {"task_id": "11", "task_code": "A1"},
            {"proj_id": "2", "task_id": "22", "task_code": "OTHER"},
        ],
        "resources": [{"rsrc_id": "1", "rsrc_name": "Renamed"}],
        "assignments": [
            {"proj_id": "1", "task_id": "11", "rsrc_id": "1"},
            {"proj_id": "2", "task_id": "22", "rsrc_id": "1"},
        ],
    }
    c, wb = workbook(tmp_path, left, right)
    assert len(c.right.assignments) == 1
    assert not c.resource_added and not c.resource_deleted
    assert not records(wb["Added Resource Assignments"])


def test_deleted_resource_assignments_preserve_duplicate_count(tmp_path):
    c, wb = workbook(tmp_path, *resource_pair())
    ws = wb["Deleted Resource Assignments"]
    rows = records(ws)
    assert [(r["Activity ID"], r["Base Resource"]) for r in rows] == [
        ("A1", "Crew"),
        ("A2", "Gone"),
    ]
    assert all(r["Revised Resource"] is None for r in rows)
    assert len(c.resource_deleted) == 2
    assert_grid(ws)


def udf_pair():
    types = [
        {
            "udf_type_id": "1",
            "udf_type_name": "udf_1",
            "udf_type_label": "Added text",
            "table_name": "TASK",
            "logical_data_type": "FT_TEXT",
        },
        {
            "udf_type_id": "2",
            "udf_type_name": "udf_2",
            "udf_type_label": "Deleted number",
            "table_name": "TASK",
            "logical_data_type": "FT_FLOAT",
        },
        {
            "udf_type_id": "3",
            "udf_type_name": "udf_3",
            "udf_type_label": "Revised date",
            "table_name": "TASK",
            "logical_data_type": "FT_DATE",
        },
        {
            "udf_type_id": "4",
            "udf_type_name": "udf_4",
            "udf_type_label": "Stable number",
            "table_name": "TASK",
            "logical_data_type": "FT_INT",
        },
        {
            "udf_type_id": "5",
            "udf_type_name": "udf_5",
            "udf_type_label": "Project value",
            "table_name": "PROJECT",
            "logical_data_type": "FT_TEXT",
        },
    ]
    left = {
        "activities": [{"task_id": "1", "task_code": "A1", "task_name": "Work"}],
        "udf_types": types,
        "udf_values": [
            {"udf_type_id": "2", "fk_id": "1", "proj_id": "1", "udf_number": "0"},
            {
                "udf_type_id": "3",
                "fk_id": "1",
                "proj_id": "1",
                "udf_date": "2024-01-01 08:00",
            },
            {"udf_type_id": "4", "fk_id": "1", "proj_id": "1", "udf_number": "1.0"},
        ],
    }
    right = {
        "activities": [{"task_id": "11", "task_code": "A1", "task_name": "Work"}],
        "udf_types": [
            {
                **t,
                "udf_type_id": str(int(t["udf_type_id"]) + 10),
                "udf_type_name": t["udf_type_name"] + "_remapped",
            }
            for t in types
        ],
        "udf_values": [
            {
                "udf_type_id": "11",
                "fk_id": "11",
                "proj_id": "1",
                "udf_text": 'Café\t"quoted"\nsecond line',
            },
            {
                "udf_type_id": "13",
                "fk_id": "11",
                "proj_id": "1",
                "udf_date": "2024-02-01 08:00",
            },
            {"udf_type_id": "14", "fk_id": "11", "proj_id": "1", "udf_number": "1"},
            {
                "udf_type_id": "15",
                "fk_id": "11",
                "proj_id": "1",
                "udf_text": "Wrong scope",
            },
            {
                "udf_type_id": "11",
                "fk_id": "11",
                "proj_id": "2",
                "udf_text": "Other project",
            },
        ],
    }
    return left, right


def test_added_udf_maps_types_task_ids_and_quoted_text(tmp_path):
    c, wb = workbook(tmp_path, *udf_pair())
    text = 'Café\t"quoted"\nsecond line'
    assert c.right.activities["A1"].udfs["Added text"] == text
    assert c.right.activities["A1"].udfs["Stable number"] == 1
    assert "Project value" not in c.right.activities["A1"].udfs
    ws = wb["Added UDF"]
    (row,) = records(ws)
    assert row["Activity ID"] == "A1"
    assert row["Revised UDF"] == "Added text"
    assert row["Revised Value"] == text
    assert row["Base UDF"] is None
    assert c.modified[0].klass == "revision"
    assert_grid(ws)


def test_deleted_udf(tmp_path):
    c, wb = workbook(tmp_path, *udf_pair())
    ws = wb["Deleted UDF"]
    (row,) = records(ws)
    assert row["Activity ID"] == "A1"
    assert row["Base UDF"] == "Deleted number"
    assert row["Base Value"] == 0
    assert row["Revised Value"] == None
    assert len(c.udf_deleted) == 1
    assert_grid(ws)


def test_revised_udf(tmp_path):
    c, wb = workbook(tmp_path, *udf_pair())
    ws = wb["Revised UDF"]
    (row,) = records(ws)
    assert row["Activity ID"] == "A1"
    assert row["Base UDF"] == "Revised date"
    assert row["Base Value"] == "2024-01-01 08:00"
    assert row["Revised Value"] == "2024-02-01 08:00"
    assert len(c.udf_modified) == 1
    assert_grid(ws)


def test_combined_is_union_without_double_counting_udfs(tmp_path):
    left, right = logic_pair()
    for side, resources, udfs in zip((left, right), resource_pair(), udf_pair()):
        side.update({k: v for k, v in resources.items() if k != "activities"})
        side.update({k: v for k, v in udfs.items() if k != "activities"})
    left["activities"][0]["task_name"] = "Old"
    right["activities"][0].update(task_name="New", status_code="TK_Active")
    left["activities"].append({"task_id": "4", "task_code": "DELETED"})
    right["activities"].append({"task_id": "14", "task_code": "ADDED"})
    c, wb = workbook(tmp_path, left, right)
    ws = wb["Changes Combined"]
    rows = records(ws)
    counts = __import__("collections").Counter(r["Change type"] for r in rows)
    assert counts == {
        "Added Task": 1,
        "Deleted Task": 1,
        "Activity Field": 2,
        "Resource Value Change": 1,
        "Added Relationship": 1,
        "Deleted Relationship": 1,
        "Lag Change": 1,
        "Added Resource Assignment": 1,
        "Deleted Resource Assignment": 2,
        "Added UDF": 1,
        "Deleted UDF": 1,
        "Revised UDF": 1,
    }
    assert {r["Activity ID"] for r in rows} == {"A1", "A2", "A3", "ADDED", "DELETED"}
    assert all(r["WBS"] == "DEMO" for r in rows)
    assert {r["Class"] for r in rows if r["Change type"] == "Activity Field"} == {
        "both"
    }
    assert_grid(ws)


def test_dcma_and_warnings_are_honest_and_visible(tmp_path):
    _, wb = workbook(
        tmp_path,
        {"project": "LEFT", "activities": [{"task_code": "A1"}]},
        {"project": "RIGHT", "activities": [{"task_code": "A1"}]},
    )
    metrics = {row["#"]: row for row in records(wb["DCMA 14"])}
    assert set(metrics) == set(range(1, 15))
    for number in (11, 12, 13, 14):
        assert metrics[number]["Left value"] == "N/A"
        assert metrics[number]["Right value"] == "N/A"
        assert metrics[number]["Left"] == metrics[number]["Right"] == "N/A"
        assert metrics[number]["Note"]
    assert any(
        "Project short names differ" in row["Warning"]
        for row in records(wb["Warnings"])
    )
    assert_grid(wb["DCMA 14"])
    assert_grid(wb["Warnings"])


REQUIRED_SHEETS = {
    "Summary",
    "Added Tasks",
    "Deleted Tasks",
    "Actual Date Change",
    "Progress Change",
    "Duration Description Change",
    "Calendar Change",
    "Constraint Change",
    "Float Change",
    "Early Start Changes",
    "Early Finish Changes",
    "Late Start Changes",
    "Late Finish Changes",
    "Variances",
    "Added Relationships",
    "Deleted Relationships",
    "Lag Change",
    "Driving Path Change",
    "Added Resource Assignments",
    "Deleted Resource Assignments",
    "Added UDF",
    "Deleted UDF",
    "Revised UDF",
    "Changes Combined",
    "DCMA 14",
    "Warnings",
}


def test_full_catalog_empty_grids(tmp_path):
    _, wb = workbook(tmp_path, {}, {})
    assert REQUIRED_SHEETS <= set(wb.sheetnames)
    for ws in wb:
        if ws.title == "Summary":
            continue
        assert_grid(ws)
        if ws.title not in {"DCMA 14", "Legend", "Quality Checks", "Check Settings"}:
            assert records(ws) == [], ws.title


def test_imported_text_is_literal_and_dates_can_be_cleared(tmp_path):
    c, wb = workbook(
        tmp_path,
        {
            "project": "=Project",
            "activities": [
                {
                    "task_code": "=A1",
                    "task_name": "Old",
                    "act_start": "2024-01-01 08:00",
                }
            ],
        },
        {
            "project": "=Project",
            "activities": [{"task_code": "=A1", "task_name": "=2+2", "act_start": ""}],
        },
    )
    assert not [
        cell.coordinate
        for ws in wb
        for row in ws
        for cell in row
        if cell.data_type == "f"
    ]
    ws = wb["Actual Date Change"]
    (row,) = records(ws)
    assert row["Activity ID"] == "=A1"
    assert row["Revised Actual start"] is None
    assert ws["E5"].fill == MOD_FILL


def test_summary_catalog_links_to_detail_counts(tmp_path):
    _, wb = workbook(
        tmp_path,
        {"activities": [{"task_code": "A1", "task_name": "Same"}]},
        {"activities": [{"task_code": "A1", "task_name": "Same"}, {"task_code": "A2"}]},
    )
    ws = wb["Summary"]
    links = {
        cell.value: (cell.hyperlink, ws.cell(cell.row, cell.column + 1).value)
        for row in ws
        for cell in row
        if cell.hyperlink
    }
    assert links["Added Tasks"][1] == 1
    assert links["Deleted Tasks"][1] == 0
    assert links["Changes Combined"][1] == 1
    assert "Added Tasks" in links["Added Tasks"][0].location
