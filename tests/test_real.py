"""Opt-in local integration coverage. Real exports are never checked in."""

from pathlib import Path

import pytest
from openpyxl import load_workbook

from xercompare.cli import main
from xercompare.model.schedule import build_schedule
from xercompare.parse.xer_parser import parse_xer
from tests.test_reports import REQUIRED_SHEETS


def test_real_xer_pair(tmp_path):
    folder = Path(__file__).parent / "fixtures" / "real"
    paths = sorted(p for p in folder.glob("*") if p.suffix.lower() == ".xer")
    if len(paths) != 2:
        pytest.skip(
            "Place exactly two local XERs in tests/fixtures/real to run this check"
        )
    schedules = [build_schedule(parse_xer(path)) for path in paths]
    schedules.sort(
        key=lambda schedule: (schedule.project.data_date, schedule.source.name)
    )
    assert (
        main([str(schedules[0].source), str(schedules[1].source), "-o", str(tmp_path)])
        == 0
    )
    wb = load_workbook(tmp_path / "compare.xlsx", read_only=True, data_only=True)
    try:
        assert REQUIRED_SHEETS <= set(wb.sheetnames)
        assert {
            "Calendar Definitions",
            "Resource Value Changes",
            "Quality Findings",
            "Milestone Slips",
        } <= set(wb.sheetnames)
    finally:
        wb.close()
    assert (tmp_path / "compare.pdf").read_bytes().startswith(b"%PDF-")
