from openpyxl import load_workbook
import pytest

from tests.fixtures.make_xer import make_xer
from xercompare.compare.engine import compare_schedules
from xercompare.gui.service import read_preview
from xercompare.model.schedule import build_schedule
from xercompare.parse.xer_parser import parse_xer
from xercompare.reports.excel import MOD_FILL, write_excel


def decimal_pair():
    types = [
        {
            "udf_type_id": str(i),
            "udf_type_label": label,
            "table_name": "TASK",
            "logical_data_type": kind,
        }
        for i, (label, kind) in enumerate(
            [
                ("Budget", "FT_MONEY"),
                ("Reserve", "FT_COST"),
                ("Allowance", "FT_COST"),
                ("Cost index", "FT_FLOAT"),
                ("Cost code", "FT_TEXT"),
            ],
            1,
        )
    ]
    sides = []
    for revised in (False, True):
        sides.append(
            dict(
                activities=[
                    {
                        "task_code": "00100",
                        "dur_hr": (10.149 if revised else 10.1234) * 8,
                        "pct": 23.4567 if revised else 12.3456,
                        "tf_hr": -9.876 if revised else -10.987,
                    },
                    {"task_code": "00200"},
                    {"task_code": "00300"},
                ],
                preds=[
                    {"pred_id": "1", "succ_id": "2", "lag": 1.2345 if revised else 0},
                    {"pred_id": "2", "succ_id": "3"},
                ],
                resources=[
                    {
                        "rsrc_id": "2" if revised else "1",
                        "rsrc_name": "New crew" if revised else "Old crew",
                    }
                ],
                assignments=[
                    {
                        "task_id": "1",
                        "rsrc_id": "2" if revised else "1",
                        "target_qty": 12.3456,
                        "remain_qty": 2.25,
                        "target_cost": 1200.4567,
                        "remain_cost": -98.7654,
                        "act_reg_cost": 0,
                    }
                ],
                udf_types=types,
                udf_values=[
                    {
                        "udf_type_id": "1",
                        "fk_id": "1",
                        "udf_number": 2345.67891 if revised else 1234.56789,
                    },
                    {
                        "udf_type_id": "3" if revised else "2",
                        "fk_id": "1",
                        "udf_number": 2.625,
                    },
                    {
                        "udf_type_id": "4",
                        "fk_id": "1",
                        "udf_number": 2.3456 if revised else 1.2345,
                    },
                    {
                        "udf_type_id": "5",
                        "fk_id": "1",
                        "udf_text": "001.567" if revised else "001.234",
                    },
                ],
            )
        )
    return sides


@pytest.fixture
def decimal_report(tmp_path):
    schedules = []
    for name, options in zip(("base", "revised"), decimal_pair()):
        path = tmp_path / f"{name}.xer"
        path.write_text(make_xer(**options), encoding="cp1252")
        schedules.append(build_schedule(parse_xer(path)))
    comparison = compare_schedules(*schedules)
    path = write_excel(comparison, tmp_path / "compare.xlsx")
    book = load_workbook(path)
    yield comparison, path, book
    book.close()


def cell_records(sheet):
    return [
        dict(zip([cell.value for cell in sheet[4]], row))
        for row in sheet.iter_rows(min_row=5)
    ]


def test_excel_decimals_keep_numeric_precision_and_detect_small_changes(decimal_report):
    comparison, _, book = decimal_report
    duration = cell_records(book["Duration Description Change"])[0]
    base, revised = (
        duration["Base Original duration (days)"],
        duration["Revised Original duration (days)"],
    )
    assert base.value == pytest.approx(10.1234)
    assert revised.value == pytest.approx(10.149)
    assert base.number_format == revised.number_format == "0.0"
    assert revised.fill == MOD_FILL  # Both display as 10.1; still an actual change.
    assert comparison.summary["modified"] == 1
    for name, side in [
        ("Added Resource Assignments", "Revised"),
        ("Deleted Resource Assignments", "Base"),
    ]:
        row = cell_records(book[name])[0]
        assert row[f"{side} Target quantity"].number_format == "0.0"
        for label in ("Target cost", "Remaining cost", "Actual regular cost"):
            assert row[f"{side} {label}"].number_format == "0.00"
            assert row[f"{side} {label}"].data_type == "n"
    for title in (
        "Added UDF",
        "Deleted UDF",
        "Revised UDF",
        "Modified",
        "Changes Combined",
    ):
        for row in cell_records(book[title]):
            name_cell = row.get("Revised UDF") or row.get("Field")
            name = name_cell.value or row["Base UDF"].value
            name = name.removeprefix("udf:")
            for side in ("Base", "Revised"):
                value = row[f"{side} Value"]
                if value.data_type == "n" and value.value is not None:
                    assert value.number_format == (
                        "0.00" if name in {"Budget", "Reserve", "Allowance"} else "0.0"
                    )
    assert book["DCMA 14"]["C5"].number_format == "0.0"
    assert book["DCMA 14"]["C5"].value == pytest.approx(100*2/3)
    assert book["DCMA 14"]["C11"].number_format in {"General", "0"}
    assert book["Summary"]["B20"].value == 3
    assert duration["Activity ID"].value == "00100"


def test_desktop_preview_matches_workbook_decimal_display(decimal_report):
    _, path, _ = decimal_report

    def rows(sheet):
        preview = read_preview(path, sheet)
        return [dict(zip(preview.headers, row)) for row in preview.rows]

    duration = rows("Duration Description Change")[0]
    assert duration["Base Original duration (days)"] == "10.1"
    assert duration["Revised Original duration (days)"] == "10.1"
    resource = rows("Added Resource Assignments")[0]
    assert resource["Revised Target quantity"] == "12.3"
    assert resource["Revised Remaining quantity"] == "2.3"
    assert resource["Revised Target cost"] == "1200.46"
    assert resource["Revised Remaining cost"] == "-98.77"
    assert resource["Revised Actual regular cost"] == "0.00"
    udfs = {row["Revised UDF"]: row for row in rows("Revised UDF")}
    assert udfs["Budget"]["Revised Value"] == "2345.68"
    assert udfs["Cost index"]["Revised Value"] == "2.3"
    assert udfs["Cost code"]["Revised Value"] == "001.567"
    assert rows("Added UDF")[0]["Revised Value"] == "2.63"
    assert duration["Activity ID"] == "00100"
    assert rows("Lag Change")[0]["Revised Lag (hours)"] == "1.2"
    assert rows("Float Change")[0]["Revised Total float (days)"] == "-1.2"
    dcma = rows("DCMA 14")
    assert dcma[0]["Left value"] == "66.7"
    assert dcma[6]["Left value"] == "1"
    assert dcma[8]["Left value"] == "3"


def test_pdf_displays_one_decimal_and_keeps_counts(
    decimal_report, tmp_path, monkeypatch
):
    from xercompare.reports import pdf

    comparison, _, _ = decimal_report
    original_table = pdf.Table
    tables = []

    def capture_table(rows, *args, **kwargs):
        tables.append(rows)
        return original_table(rows, *args, **kwargs)

    monkeypatch.setattr(pdf, "Table", capture_table)
    result = pdf.write_pdf(comparison, tmp_path / "compare.pdf")
    dcma = next(rows for rows in tables if rows[0][0] == "#")
    assert dcma[1][2:4] == ["66.7", "66.7"]
    assert dcma[7][2:4] == ["1", "1"]
    assert dcma[9][2:4] == ["3", "3"]
    assert result.read_bytes().startswith(b"%PDF-")
