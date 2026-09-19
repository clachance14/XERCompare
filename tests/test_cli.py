from pathlib import Path

from xercompare.cli import main
from tests.fixtures.make_xer import make_xer


def test_cli(tmp_path: Path):
    left = tmp_path / "left.xer"
    right = tmp_path / "right.xer"
    left.write_text(
        make_xer(activities=[{"task_code": "A1", "task_name": "One"}]),
        encoding="cp1252",
    )
    right.write_text(
        make_xer(
            activities=[
                {"task_code": "A1", "task_name": "Two"},
                {"task_code": "A2", "task_name": "New"},
            ]
        ),
        encoding="cp1252",
    )
    out = tmp_path / "out"
    rc = main([str(left), str(right), "-o", str(out)])
    assert rc == 0
    assert (out / "compare.xlsx").exists()
    assert (out / "compare.pdf").exists()


def test_cli_offline_with_markup_in_project_names(tmp_path, monkeypatch):
    import socket
    from openpyxl import load_workbook
    from tests.test_reports import REQUIRED_SHEETS

    def no_network(*args, **kwargs):
        raise AssertionError("The comparison must remain offline")

    monkeypatch.setattr(socket.socket, "connect", no_network)
    monkeypatch.setattr(socket, "create_connection", no_network)
    left, right = tmp_path / "left.xer", tmp_path / "right.xer"
    left.write_text(
        make_xer(project="North <b> & café", activities=[{"task_code": "A1"}]),
        encoding="cp1252",
    )
    right.write_text(
        make_xer(
            project="South <b> & café", activities=[{"task_code": "A1", "pct": 10}]
        ),
        encoding="cp1252",
    )
    out = tmp_path / "reports"
    assert main([str(left), str(right), "-o", str(out)]) == 0
    wb = load_workbook(out / "compare.xlsx")
    assert REQUIRED_SHEETS <= set(wb.sheetnames)
    assert (out / "compare.pdf").read_bytes().startswith(b"%PDF-")


def test_cli_reports_resource_only_changes(tmp_path, capsys):
    left, right = tmp_path / "left.xer", tmp_path / "right.xer"
    common = {
        "activities": [{"task_id": "1", "task_code": "A1"}],
        "resources": [{"rsrc_id": "1", "rsrc_name": "Crew"}],
    }
    left.write_text(make_xer(**common), encoding="cp1252")
    right.write_text(
        make_xer(**common, assignments=[{"task_id": "1", "rsrc_id": "1"}]),
        encoding="cp1252",
    )
    assert main([str(left), str(right), "-o", str(tmp_path / "out")]) == 0
    assert "resource +1 / -0" in capsys.readouterr().out


def test_cli_cp1252_console(tmp_path, monkeypatch):
    from io import BytesIO, TextIOWrapper
    import sys

    left, right = tmp_path / "left.xer", tmp_path / "right.xer"
    for path in (left, right):
        path.write_text(make_xer(activities=[{"task_code": "A1"}]), encoding="cp1252")
    buffer = BytesIO()
    console = TextIOWrapper(
        buffer, encoding="cp1252", errors="strict", write_through=True
    )
    monkeypatch.setattr(sys, "stdout", console)
    assert main([str(left), str(right), "-o", str(tmp_path / "out")]) == 0
    assert b"left.xer (1 act) -> right.xer (1 act)" in buffer.getvalue()
