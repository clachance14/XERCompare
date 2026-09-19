"""Guard the real-export formatting regression without a machine-speed timeout."""

from tests.test_reports import schedule
from xercompare.compare.engine import compare_schedules
from xercompare.reports import excel
from openpyxl import load_workbook


def test_formatting_setup_does_not_grow_per_activity_cell(tmp_path, monkeypatch):
    original = excel.Alignment
    calls = []

    def counted(*args, **kwargs):
        calls.append(1)
        return original(*args, **kwargs)

    monkeypatch.setattr(excel, "Alignment", counted)
    counts = []
    for size in (3, 30):
        left = schedule(
            tmp_path / "l.xer",
            activities=[
                {"task_code": str(i), "task_name": "Before"} for i in range(size)
            ],
        )
        right = schedule(
            tmp_path / "r.xer",
            activities=[
                {"task_code": str(i), "task_name": "After"} for i in range(size)
            ],
        )
        calls.clear()
        path = excel.write_excel(
            compare_schedules(left, right), tmp_path / f"{size}.xlsx"
        )
        counts.append(len(calls))
        wb = load_workbook(path)
        ws = wb["Duration Description Change"]
        assert (
            ws["G5"].fill == excel.MOD_FILL
        )  # Changed description; paired numeric fields retain their own style.
        assert ws["D5"].number_format == "0.0"
        assert ws["D5"].fill != excel.MOD_FILL
    assert counts[1] <= counts[0] + 10, counts
