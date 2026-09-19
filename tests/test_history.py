from threading import Event
from pathlib import Path
import pytest
from openpyxl import load_workbook
from tests.fixtures.make_xer import make_xer
from xercompare.history import discover_history, HistoryRequest, run_history
from xercompare.gui.service import Cancelled


def snapshots(folder):
    folder.mkdir()
    for filename, date, count in [
        ("z.xer", "2024-01-01 00:00", 1),
        ("a.xer", "2024-02-01 00:00", 2),
        ("b.XER", "2024-02-01 00:00", 3),
    ]:
        (folder / filename).write_text(
            make_xer(
                data_date=date,
                activities=[{"task_code": f"A{i}"} for i in range(count)],
            ),
            encoding="cp1252",
        )
    return discover_history(folder)


def test_history_order_trends_full_pair_reports_baseline(tmp_path):
    items = snapshots(tmp_path / "inputs")
    assert [s.path.name for s in items] == ["z.xer", "a.xer", "b.XER"]
    request = HistoryRequest(
        tuple(s.path for s in items), tmp_path / "out", baseline=items[0].path
    )
    result = run_history(request)
    wb = load_workbook(result.excel)
    assert [r[3] for r in wb["Trends"].iter_rows(min_row=5, values_only=True)] == [
        1,
        2,
        3,
    ]
    assert (
        len(result.pairs) == 4
    )  # 2 consecutive, 2 baseline vs update (skip baseline itself)
    assert all(
        (result.folder / p.report / "compare.xlsx").exists() for p in result.pairs
    )
    assert any("same data date" in w.lower() for w in result.warnings)
    assert len(wb["Trends"]._charts) == 4
    assert result.pdf.read_bytes().startswith(b"%PDF-")


def test_history_rejects_mixed_project_without_acknowledgment(tmp_path):
    items = snapshots(tmp_path / "inputs")
    items[1].path.write_text(make_xer(project="OTHER"), encoding="cp1252")
    request = HistoryRequest(tuple(s.path for s in items), tmp_path / "out")
    with pytest.raises(ValueError, match="project"):
        run_history(request)
    assert not request.output.exists()


def test_history_cancellation_does_not_publish_partial_results(tmp_path):
    items = snapshots(tmp_path / "inputs")
    cancel = Event()

    def progress(value, message):
        if value >= 30:
            cancel.set()

    with pytest.raises(Cancelled):
        run_history(
            HistoryRequest(tuple(s.path for s in items), tmp_path / "out"),
            progress,
            cancel,
        )
    assert not list((tmp_path / "out").glob("XERHistory-*"))
    assert not list((tmp_path / "out").glob(".xerhistory-*"))


def test_history_requires_valid_dates_and_distinct_inputs(tmp_path):
    folder = tmp_path / "inputs"
    folder.mkdir()
    p = folder / "bad.xer"
    p.write_text(make_xer(data_date=""), encoding="cp1252")
    with pytest.raises(ValueError, match="data date"):
        discover_history(folder)
    with pytest.raises(ValueError, match="distinct"):
        HistoryRequest((p, p), tmp_path / "out").validate()
