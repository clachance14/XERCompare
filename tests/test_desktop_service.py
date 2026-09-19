from pathlib import Path
from threading import Event

import pytest
from openpyxl import load_workbook

from tests.fixtures.make_xer import make_xer
from xercompare.gui.service import (
    Cancelled,
    CompareRequest,
    run_comparison,
    read_preview,
)


def request(tmp_path):
    left, right = tmp_path / "base.xer", tmp_path / "revised.xer"
    left.write_text(
        make_xer(activities=[{"task_code": "A1", "task_name": "Old"}]),
        encoding="cp1252",
    )
    right.write_text(
        make_xer(
            activities=[{"task_code": "A1", "task_name": "New"}, {"task_code": "A2"}]
        ),
        encoding="cp1252",
    )
    return CompareRequest(left, right, tmp_path / "reports")


def test_desktop_run_and_report_preview(tmp_path):
    req = request(tmp_path)
    progress = []
    result = run_comparison(req, lambda value, text: progress.append((value, text)))
    assert result.folder.parent == req.output
    assert result.summary["added"] == 1
    assert result.summary["modified"] == 1
    assert len(result.reports) == 39
    assert result.pdf.read_bytes().startswith(b"%PDF-")
    assert load_workbook(result.excel)["Added Tasks"]["A5"].value == "A2"
    preview = read_preview(result.excel, "Added Tasks")
    assert preview.headers[:3] == ["Activity ID", "Name", "WBS"]
    assert preview.rows[0][0] == "A2"
    assert progress[-1][0] == 100
    assert all(a[0] <= b[0] for a, b in zip(progress, progress[1:]))


def test_desktop_rerun_preserves_previous_reports(tmp_path):
    req = request(tmp_path)
    first = run_comparison(req)
    original = first.excel.read_bytes()
    second = run_comparison(req)
    assert first.folder != second.folder
    assert first.excel.read_bytes() == original
    assert second.excel.exists()


def test_failed_report_does_not_publish_partial_results(tmp_path, monkeypatch):
    from xercompare.gui import service

    req = request(tmp_path)

    def broken_pdf(*args):
        raise OSError("Disk full")

    monkeypatch.setattr(service, "write_pdf", broken_pdf)
    with pytest.raises(OSError, match="Disk full"):
        run_comparison(req)
    assert not list(req.output.iterdir())


def test_cancelled_run_cleans_staged_reports(tmp_path):
    req = request(tmp_path)
    cancel = Event()

    def progress(value, text):
        if value >= 80:
            cancel.set()

    with pytest.raises(Cancelled):
        run_comparison(req, progress, cancel)
    assert not list(req.output.iterdir())


def test_missing_input_is_actionable_and_creates_no_output(tmp_path):
    req = request(tmp_path)
    req.left.unlink()
    with pytest.raises(ValueError, match="Older XER"):
        run_comparison(req)
    assert not req.output.exists()


def test_preview_is_bounded_without_truncating_export(tmp_path):
    req = request(tmp_path)
    req.right.write_text(
        make_xer(activities=[{"task_code": f"B{i}"} for i in range(210)]),
        encoding="cp1252",
    )
    result = run_comparison(req)
    preview = read_preview(result.excel, "Added Tasks", limit=200)
    assert len(preview.rows) == 200
    assert preview.total == 210
    assert load_workbook(result.excel)["Added Tasks"].max_row == 214
