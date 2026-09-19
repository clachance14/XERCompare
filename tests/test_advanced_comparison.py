from openpyxl import load_workbook
from tests.fixtures.make_xer import calendar_data
from tests.test_reports import schedule, records, assert_grid
from xercompare.compare.engine import compare_schedules
from xercompare.reports.excel import write_excel
from xercompare.model.calendars import resolve_calendar


def test_calendar_definition_inheritance_and_normalization(tmp_path):
    def make(name, finish, order=False):
        blob = calendar_data({2: [("8:00", finish)], 1: []}, {45292: []})
        if order:
            blob = blob.replace("s|8:00|f|17:00", "f|17:00|s|08:00").replace(
                ")(", ")\x7f ("
            )
        return schedule(
            tmp_path / name,
            activities=[{"task_code": "A", "clndr_id": "2"}],
            calendars=[
                {
                    "clndr_id": "1",
                    "clndr_name": "Parent",
                    "day_hr_cnt": 8,
                    "clndr_data": blob,
                },
                {
                    "clndr_id": "2",
                    "clndr_name": "Child",
                    "base_clndr_id": "1",
                    "day_hr_cnt": 8,
                    "clndr_data": calendar_data({3: [("19:00", "00:00")]}),
                },
            ],
        )

    left, same, right = (
        make("l.xer", "17:00"),
        make("s.xer", "17:00", True),
        make("r.xer", "18:00"),
    )
    definition = resolve_calendar(left, "2")
    assert definition.week[2] == (("08:00", "17:00"),)
    assert definition.week[3] == (("19:00", "24:00"),)
    assert definition.exceptions["2024-01-01"] == ()
    assert not compare_schedules(left, same).calendar_changes
    c = compare_schedules(left, right)
    assert any(
        d.task_code == "A" and "Monday" in d.field and "18:00" in d.after
        for d in c.calendar_changes
    )
    wb = load_workbook(write_excel(c, tmp_path / "c.xlsx"))
    assert_grid(wb["Calendar Definitions"])
    assert any(r["Activity ID"] == "A" for r in records(wb["Calendar Definitions"]))


def test_bad_calendar_is_unknown_not_equal_to_nonworking(tmp_path):
    s = schedule(
        tmp_path / "a.xer",
        calendars=[
            {"clndr_id": "1", "clndr_name": "Bad", "clndr_data": "broken"},
            {
                "clndr_id": "2",
                "clndr_name": "Cycle",
                "base_clndr_id": "2",
                "clndr_data": calendar_data(),
            },
        ],
    )
    assert not resolve_calendar(s, "1").available
    assert "parse" in resolve_calendar(s, "1").note.lower()
    assert not resolve_calendar(s, "2").available


def test_matched_resource_changes_with_reordered_duplicates(tmp_path):
    def make(name, rows):
        return schedule(
            tmp_path / name,
            activities=[{"task_code": "A"}],
            resources=[{"rsrc_id": "7", "rsrc_name": "Crew"}],
            assignments=[{"task_id": "1", "rsrc_id": "7", **r} for r in rows],
        )

    left = make(
        "l.xer",
        [
            {
                "guid": "a",
                "target_qty": 10,
                "target_cost": 100.123,
                "target_qty_per_hr": 2,
            },
            {"guid": "b", "target_qty": 20, "target_cost": 200},
        ],
    )
    right = make(
        "r.xer",
        [
            {"guid": "b", "target_qty": 20, "target_cost": 200},
            {
                "guid": "a",
                "target_qty": 12,
                "target_cost": 130.456,
                "target_qty_per_hr": 3,
            },
        ],
    )
    c = compare_schedules(left, right)
    assert not c.resource_added and not c.resource_deleted
    assert len(c.resource_modified) == 1
    assert {f.field for f in c.resource_modified[0].changes} == {
        "target_qty",
        "target_cost",
        "target_qty_per_hr",
    }
    wb = load_workbook(write_excel(c, tmp_path / "c.xlsx"))
    ws = wb["Resource Value Changes"]
    assert_grid(ws)
    rows = records(ws)
    costrow = next(i for i, r in enumerate(rows, 5) if r["Field"] == "Target cost")
    assert ws.cell(costrow, 5).value == 130.456
    assert ws.cell(costrow, 5).number_format == "0.00"
    assert (
        len(
            [
                r
                for r in records(wb["Changes Combined"])
                if r["Change type"] == "Resource Value Change"
            ]
        )
        == 3
    )


def test_calendar_exception_added_and_removed(tmp_path):
    def make(name, exceptions):
        return schedule(
            tmp_path / name,
            activities=[{"task_code": "A"}],
            calendars=[
                {
                    "clndr_id": "1",
                    "clndr_name": "Week",
                    "clndr_data": calendar_data({2: [("08:00", "17:00")]}, exceptions),
                }
            ],
        )

    c = compare_schedules(
        make("l.xer", {45292: []}), make("r.xer", {45293: [("08:00", "12:00")]})
    )
    changes = {d.field: d for d in c.calendar_changes}
    assert changes["Exception 2024-01-01"].before == "Nonworking"
    assert changes["Exception 2024-01-01"].after is None
    assert changes["Exception 2024-01-02"].before is None
    assert changes["Exception 2024-01-02"].after == "08:00-12:00"


def test_assignment_ambiguity_warns_without_losing_multiplicity(tmp_path):
    def make(name, quantities):
        return schedule(
            tmp_path / name,
            activities=[{"task_code": "A"}],
            resources=[{"rsrc_id": "7", "rsrc_name": "Crew"}],
            assignments=[
                {"task_id": "1", "rsrc_id": "7", "target_qty": q} for q in quantities
            ],
        )

    c = compare_schedules(make("l.xer", [10, 20]), make("r.xer", [11, 21]))
    assert (
        len(c.resource_modified) == 2
        and not c.resource_added
        and not c.resource_deleted
    )
    assert any("ambiguous" in w for w in c.warnings)


def test_visible_code_assignment_addition_marks_all_affected_tasks(tmp_path):
    tasks = [{"task_code": "A"}, {"task_code": "B"}]
    left = schedule(tmp_path / "l.xer", activities=tasks)
    right = schedule(
        tmp_path / "r.xer",
        activities=tasks,
        tables={
            "ACTVTYPE": [{"actv_code_type_id": "9", "actv_code_type": "Level Order"}],
            "ACTVCODE": [
                {"actv_code_type_id": "9", "actv_code_id": "90", "short_name": "2"}
            ],
            "TASKACTV": [
                {"task_id": "1", "actv_code_id": "90"},
                {"task_id": "2", "actv_code_id": "90"},
            ],
        },
    )
    c = compare_schedules(left, right)
    assert len(c.modified) == 2
    assert all(
        d.changes[0].field == "code:Level Order" and d.klass == "revision"
        for d in c.modified
    )
