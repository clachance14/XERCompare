"""Multi-sheet Excel comparison workbook."""

from __future__ import annotations

from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.worksheet.hyperlink import Hyperlink

from xercompare.compare.engine import Comparison, _norm, RESOURCE_FIELDS
from xercompare.quality.dcma import assess
from xercompare.model.schedule import Activity
from xercompare.reports.formatting import COST_FORMAT, NUMBER_FORMAT


NAVY = "1B365D"
NAVY_FILL = PatternFill("solid", fgColor=NAVY)
STEEL = "2E75B6"
STEEL_FILL = PatternFill("solid", fgColor=STEEL)
HEADER_FONT = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(name="Calibri", color="FFFFFF", bold=True, size=18)
SUB_FONT = Font(name="Calibri", color="D6E3F0", size=11)
LABEL_FONT = Font(name="Calibri", bold=True, size=11, color="1B365D")
BODY_FONT = Font(name="Calibri", size=11)
WHITE_FONT = Font(name="Calibri", color="FFFFFF", bold=True, size=11)

ADDED_FILL = PatternFill("solid", fgColor="C6EFCE")
DELETED_FILL = PatternFill("solid", fgColor="FFC7CE")
MOD_FILL = PatternFill("solid", fgColor="FFE699")
PROGRESS_FILL = PatternFill("solid", fgColor="BDD7EE")
REVISION_FILL = PatternFill("solid", fgColor="F8CBAD")
DATE_FILL = PatternFill("solid", fgColor="E2D5F1")
BOTH_FILL = PatternFill("solid", fgColor="F4B183")
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")
NA_FILL = PatternFill("solid", fgColor="D9D9D9")
ZEBRA = PatternFill("solid", fgColor="F2F2F2")
LIGHT_NAVY = PatternFill("solid", fgColor="D6E3F0")
WHITE = PatternFill("solid", fgColor="FFFFFF")
KPI_FILL = PatternFill("solid", fgColor="0D1B2A")

THIN = Border(
    left=Side(style="thin", color="BDD7EE"),
    right=Side(style="thin", color="BDD7EE"),
    top=Side(style="thin", color="BDD7EE"),
    bottom=Side(style="thin", color="BDD7EE"),
)

CLASS_FILL = {
    "progress": PROGRESS_FILL,
    "revision": REVISION_FILL,
    "date": DATE_FILL,
    "both": BOTH_FILL,
    "n/a": ZEBRA,
}


# Labels are shared by the dedicated sheets and the packed variance view.
FIELD_SHEETS = {
    "Actual Date Change": [("act_start", "Actual start"), ("act_end", "Actual finish")],
    "Progress Change": [
        ("remain_dur_days", "Remaining duration (days)"),
        ("phys_complete_pct", "Physical complete (%)"),
        ("status_code", "Status"),
        ("restart", "Restart"),
        ("reend", "Re-end"),
        ("suspend", "Suspend"),
        ("resume", "Resume"),
    ],
    "Duration Description Change": [
        ("orig_dur_days", "Original duration (days)"),
        ("task_name", "Description"),
    ],
    "Calendar Change": [("calendar_name", "Calendar")],
    "Constraint Change": [
        ("cstr_type", "Constraint type"),
        ("cstr_date", "Constraint date"),
        ("cstr_type2", "Secondary constraint type"),
        ("cstr_date2", "Secondary constraint date"),
    ],
    "Float Change": [
        ("total_float_days", "Total float (days)"),
        ("free_float_days", "Free float (days)"),
    ],
    "Early Start Changes": [("early_start", "Early start")],
    "Early Finish Changes": [("early_end", "Early finish")],
    "Late Start Changes": [("late_start", "Late start")],
    "Late Finish Changes": [("late_end", "Late finish")],
}


def _sheet_field_changes(wb, c: Comparison, title, fields, note="") -> None:
    rows = []
    watched = {field for field, _ in fields}
    for delta in c.modified:
        if not watched.intersection(change.field for change in delta.changes):
            continue
        left, right = (
            c.left.activities[delta.task_code],
            c.right.activities[delta.task_code],
        )
        pairs = [(getattr(left, field), getattr(right, field)) for field, _ in fields]
        rows.append((right, pairs, [delta.klass]))
    _paired_sheet(wb, title, [label for _, label in fields], rows, ["Class"], note)


def write_excel(comparison: Comparison, dest: str | Path) -> Path:
    dest = Path(dest)
    wb = Workbook()
    wb.properties.title = "XERCompare schedule comparison"
    wb.properties.creator = "XERCompare"

    _summary(wb.active, comparison)
    _sheet_activities(wb, "Added Tasks", comparison, added=True)
    _sheet_activities(wb, "Deleted Tasks", comparison, added=False)
    for title, fields in FIELD_SHEETS.items():
        _sheet_field_changes(wb, comparison, title, fields)
    _sheet_date_variance(wb, comparison)
    _sheet_relationships(wb, comparison, "Added Relationships", comparison.logic_added)
    _sheet_relationships(
        wb, comparison, "Deleted Relationships", comparison.logic_deleted
    )
    _sheet_relationships(wb, comparison, "Lag Change", comparison.logic_modified)
    _sheet_driving_path(wb, comparison)
    _sheet_resources(
        wb, comparison, "Added Resource Assignments", comparison.resource_added, True
    )
    _sheet_resources(
        wb,
        comparison,
        "Deleted Resource Assignments",
        comparison.resource_deleted,
        False,
    )
    _sheet_udfs(wb, comparison, "Added UDF", comparison.udf_added)
    _sheet_udfs(wb, comparison, "Deleted UDF", comparison.udf_deleted)
    _sheet_udfs(wb, comparison, "Revised UDF", comparison.udf_modified)
    _sheet_advanced_values(wb, comparison)
    _sheet_combined(wb, comparison)
    _sheet_dcma(wb, comparison)
    from xercompare.reports.quality import quality_sheets

    quality_sheets(wb, comparison)
    _sheet_warnings(wb, comparison)
    _sheet_modified(wb, comparison)
    _sheet_logic(wb, comparison)
    _cover_legend(wb)
    _summary_catalog(wb)

    for ws in wb.worksheets:
        # These reports are stored-value snapshots; exported strings must never
        # become executable Excel formulas, including file/project identities.
        for row in ws:
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
                if isinstance(cell.value, float) and cell.number_format == "General":
                    cell.number_format = NUMBER_FORMAT
        if ws.auto_filter.ref:
            ws.auto_filter.ref = (
                f"A4:{get_column_letter(ws.max_column)}{max(4, ws.max_row)}"
            )

    dest.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest)
    return dest


def _banner(ws: Worksheet, subtitle: str, last_col: int = 8) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    a1 = ws["A1"]
    a1.value = "XERCOMPARE"
    a1.font = TITLE_FONT
    a1.fill = NAVY_FILL
    a1.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    a2 = ws["A2"]
    a2.value = subtitle
    a2.font = SUB_FONT
    a2.fill = STEEL_FILL
    a2.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    for col in range(1, last_col + 1):
        ws.cell(1, col).fill = NAVY_FILL
        ws.cell(2, col).fill = STEEL_FILL


def _style_header_row(ws: Worksheet, headers: list[str], row: int = 4) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row, col, name)
        cell.fill = NAVY_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = THIN
    ws.freeze_panes = f"A{row + 1}"
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(headers))}{row}"
    ws.row_dimensions[row].height = 34
    widths = {
        "Activity ID": 16,
        "Name": 42,
        "Base Description": 42,
        "Revised Description": 42,
        "Name (new)": 36,
        "Name (old)": 36,
        "Class": 12,
        "Field": 22,
        "Before": 28,
        "After": 28,
        "Status": 12,
        "Predecessor": 16,
        "Successor": 16,
        "Type": 10,
        "Metric": 32,
        "Warning": 90,
        "WBS": 28,
    }
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(
            h, min(max(len(h) + 3, 14), 28)
        )


def _paint_row(ws: Worksheet, row: int, cols: int, fill: PatternFill) -> None:
    # Register each body style once per workbook. Rebuilding and hashing four
    # style objects for every cell dominated exports of real schedules.
    book = ws.parent
    if not hasattr(book, "_xercompare_body_styles"):
        book._xercompare_body_styles = {}
    cache = book._xercompare_body_styles
    key = id(fill)
    if key not in cache:
        style = NamedStyle(
            name=f"XERCompare body {len(cache)}",
            font=BODY_FONT,
            fill=fill,
            border=THIN,
            alignment=Alignment(vertical="center", wrap_text=True),
        )
        book.add_named_style(style)
        cache[key] = style._style
    for col in range(1, cols + 1):
        cell = ws.cell(row, col)
        number_format_id = cell._style.numFmtId if cell.has_style else 0
        # Independent arrays keep later revised-cell highlights from changing
        # neighboring cells. Preserve each cell's numeric/currency format.
        cell._style = copy(cache[key])
        cell._style.numFmtId = number_format_id


def _summary(ws: Worksheet, c: Comparison) -> None:
    ws.title = "Summary"
    _banner(
        ws,
        "Offline Primavera P6 schedule comparison  ·  stored-field delta  ·  not a CPM reschedule",
        6,
    )
    ws.sheet_view.showGridLines = False

    # File identity
    ws["A4"] = "FILES"
    ws["A4"].font = WHITE_FONT
    ws["A4"].fill = NAVY_FILL
    ws.merge_cells("A4:B4")

    identity = [
        ("Left file (older / baseline)", c.left.source.name),
        ("Left full path", str(c.left.source)),
        ("Left project", c.left.project.short_name),
        ("Left data date", c.left.project.data_date or "—"),
        ("Left P6 export version", c.left.p6_version or "—"),
        ("Right file (newer / update)", c.right.source.name),
        ("Right full path", str(c.right.source)),
        ("Right project", c.right.project.short_name),
        ("Right data date", c.right.project.data_date or "—"),
        ("Right P6 export version", c.right.p6_version or "—"),
        ("Report generated", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for i, (label, value) in enumerate(identity, start=5):
        ws.cell(i, 1, label).font = LABEL_FONT
        ws.cell(i, 1).fill = LIGHT_NAVY
        ws.cell(i, 2, value).font = BODY_FONT

    # KPI tiles
    s = c.summary
    klasses = Counter(d.klass for d in c.modified)
    tiles = [
        ("LEFT ACT", s["left_activities"], NAVY_FILL),
        ("RIGHT ACT", s["right_activities"], NAVY_FILL),
        ("ADDED", s["added"], ADDED_FILL),
        ("DELETED", s["deleted"], DELETED_FILL),
        ("MODIFIED", s["modified"], MOD_FILL),
        ("UNCHANGED", s["unchanged"], ZEBRA),
        ("LOGIC +", s["logic_added"], ADDED_FILL),
        ("LOGIC −", s["logic_deleted"], DELETED_FILL),
        ("LAG Δ", s["logic_modified"], MOD_FILL),
        ("PROGRESS", klasses.get("progress", 0), PROGRESS_FILL),
        ("REVISION", klasses.get("revision", 0), REVISION_FILL),
        ("BOTH / DATE", klasses.get("both", 0) + klasses.get("date", 0), BOTH_FILL),
    ]
    start = 18
    ws.cell(start, 1, "COUNTS").font = WHITE_FONT
    ws.cell(start, 1).fill = NAVY_FILL
    ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=6)
    headers = ["Metric", "Count", "Metric", "Count", "Metric", "Count"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(start + 1, col, h)
        cell.fill = STEEL_FILL
        cell.font = HEADER_FONT
    for i, (label, value, fill) in enumerate(tiles):
        col_block = (i % 3) * 2
        row = start + 2 + (i // 3)
        ws.cell(row, col_block + 1, label).fill = fill
        ws.cell(row, col_block + 1).font = (
            WHITE_FONT if fill == NAVY_FILL else LABEL_FONT
        )
        ws.cell(row, col_block + 1).border = THIN
        ws.cell(row, col_block + 2, value).fill = fill
        ws.cell(row, col_block + 2).font = Font(
            name="Calibri",
            bold=True,
            size=14,
            color="FFFFFF" if fill == NAVY_FILL else NAVY,
        )
        ws.cell(row, col_block + 2).border = THIN
        ws.cell(row, col_block + 2).alignment = Alignment(horizontal="center")

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 55
    for col in "CDEF":
        ws.column_dimensions[col].width = 16

    note_row = start + 8
    ws.cell(note_row, 1, "HOW TO READ").font = WHITE_FONT
    ws.cell(note_row, 1).fill = NAVY_FILL
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=6)
    notes = [
        "Green = added in the update. Red = removed since the left file. Gold = field changed.",
        "Progress = actuals / remaining duration / percent complete / status. Revision = name, original duration, calendar, constraints, WBS, codes, logic.",
        "Date = early/late/target dates and total float — effects of the last P6 schedule run, not proof of a planner edit.",
        "This workbook compares stored XER fields. It does not recalculate CPM. Float and criticality are only as good as the last schedule of each file.",
    ]
    for i, text in enumerate(notes):
        ws.merge_cells(
            start_row=note_row + 1 + i,
            start_column=1,
            end_row=note_row + 1 + i,
            end_column=6,
        )
        ws.cell(note_row + 1 + i, 1, text).font = BODY_FONT
        ws.cell(note_row + 1 + i, 1).alignment = Alignment(
            wrap_text=True, vertical="center"
        )
        ws.row_dimensions[note_row + 1 + i].height = 30


def _cover_legend(wb: Workbook) -> None:
    ws = wb.create_sheet("Legend")
    _banner(ws, "Color codes and field classification", 4)
    _style_header_row(
        ws, ["Class / color", "Meaning", "Typical fields", "Review use"], row=4
    )
    rows = [
        (
            "Added",
            "Activity ID exists only in the right file",
            "new task_code",
            "Scope growth, re-IDs, forgotten deletes",
            ADDED_FILL,
        ),
        (
            "Deleted",
            "Activity ID exists only in the left file",
            "missing task_code",
            "Scope cut or ID rename (looks like delete+add)",
            DELETED_FILL,
        ),
        (
            "Progress",
            "Update movement you expect from time passing",
            "act_start, act_end, remain_dur, status, %",
            "Period progress check",
            PROGRESS_FILL,
        ),
        (
            "Revision",
            "Planner or network edit",
            "name, orig dur, calendar, constraint, WBS, codes",
            "Unauthorized change / MIP 3.4 later",
            REVISION_FILL,
        ),
        (
            "Date",
            "Stored date or float moved",
            "ES/EF/LS/LF/target, total float",
            "Slippage list — cause may be elsewhere",
            DATE_FILL,
        ),
        (
            "Both",
            "Progress and revision on the same activity",
            "mixed",
            "Do not treat as clean progress",
            BOTH_FILL,
        ),
        (
            "DCMA PASS",
            "Metric under conventional threshold",
            "see DCMA 14 sheet",
            "Quality both sides",
            PASS_FILL,
        ),
        (
            "DCMA FAIL",
            "Metric over threshold",
            "see DCMA 14 sheet",
            "Owner review flag",
            FAIL_FILL,
        ),
        (
            "DCMA N/A",
            "Needs CPM engine or a baseline pair",
            "CPLI, BEI, CP test",
            "Do not invent a number",
            NA_FILL,
        ),
    ]
    for i, (a, b, c, d, fill) in enumerate(rows, start=5):
        ws.cell(i, 1, a)
        ws.cell(i, 2, b)
        ws.cell(i, 3, c)
        ws.cell(i, 4, d)
        _paint_row(ws, i, 4, fill)
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 42
    ws.column_dimensions["D"].width = 42
    ws.sheet_view.showGridLines = False


def _paired_sheet(wb, title, fields, rows, extra_headers=(), note="") -> Worksheet:
    """A row is (activity context, base/revised pairs, trailing context cells)."""
    ws = wb.create_sheet(title)
    headers = ["Activity ID", "Name", "WBS"]
    for label in fields:
        headers.extend([f"Base {label}", f"Revised {label}"])
    headers.extend(extra_headers)
    _banner(
        ws,
        f"{title}  ·  {len(rows)} row(s)" + (f"  ·  {note}" if note else ""),
        len(headers),
    )
    _style_header_row(ws, headers)
    for row, (activity, pairs, context) in enumerate(rows, start=5):
        values = [activity.task_code, activity.task_name, activity.wbs_path]
        values.extend(value for pair in pairs for value in pair)
        values.extend(context)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row, col, value)
            # Imported text is always literal, even when it begins with '='.
            if isinstance(value, str):
                cell.data_type = "s"
            if isinstance(value, (int, float)) and not isinstance(value, bool):
                cell.number_format = (
                    COST_FORMAT if "cost" in headers[col - 1].lower() else NUMBER_FORMAT
                )
        _paint_row(ws, row, len(headers), ZEBRA if row % 2 else WHITE)
        if title.startswith("Added "):
            ws.cell(row, 1).fill = ADDED_FILL
        elif title.startswith("Deleted "):
            ws.cell(row, 1).fill = DELETED_FILL
        for index, label in enumerate(extra_headers, start=4 + 2 * len(fields)):
            value = ws.cell(row, index).value
            if label == "Class":
                ws.cell(row, index).fill = CLASS_FILL.get(value, ZEBRA)
            elif label == "Change type" and isinstance(value, str):
                if value.startswith("Added "):
                    ws.cell(row, index).fill = ADDED_FILL
                elif value.startswith("Deleted "):
                    ws.cell(row, index).fill = DELETED_FILL
        for index, (before, after) in enumerate(pairs):
            if _norm(before) != _norm(after):
                ws.cell(row, 5 + index * 2).fill = MOD_FILL
    return ws


def _format_udf_rows(ws, comparison, entries, base_column=4):
    """Use each side's exported UDF type, including money with arbitrary labels."""
    for row, code, name in entries:
        for column, schedule in [
            (base_column, comparison.left),
            (base_column + 1, comparison.right),
        ]:
            activity = schedule.activities.get(code)
            if activity and activity.udf_types.get(name) in {"FT_MONEY", "FT_COST"}:
                ws.cell(row, column).number_format = COST_FORMAT


def _sheet_activities(wb, title, c: Comparison, added: bool) -> None:
    schedule = c.right if added else c.left
    items = c.added if added else c.deleted
    rows = []
    for delta in items:
        activity = schedule.activities[delta.task_code]
        pair = (None, delta.task_code) if added else (delta.task_code, None)
        rows.append((activity, [pair], [delta.klass]))
    _paired_sheet(wb, title, ["Activity ID"], rows, ["Class"])


def _sheet_modified(wb, c: Comparison) -> None:
    rows = []
    for delta in c.modified:
        for change in delta.changes:
            rows.append(
                (
                    c.right.activities[delta.task_code],
                    [(change.before, change.after)],
                    [delta.klass, change.field, change.klass],
                )
            )
    ws = _paired_sheet(
        wb,
        "Modified",
        ["Value"],
        rows,
        ["Class", "Field", "Field class"],
        "Raw activity field changes",
    )
    _format_udf_rows(
        ws,
        c,
        (
            (row, activity.task_code, context[1].removeprefix("udf:"))
            for row, (activity, _, context) in enumerate(rows, 5)
            if context[1].startswith("udf:")
        ),
    )


def _sheet_date_variance(wb, c: Comparison) -> None:
    fields = [
        ("early_start", "Early start"),
        ("early_end", "Early finish"),
        ("late_start", "Late start"),
        ("late_end", "Late finish"),
        ("total_float_days", "Total float (days)"),
        ("free_float_days", "Free float (days)"),
        ("orig_dur_days", "Original duration (days)"),
        ("remain_dur_days", "Remaining duration (days)"),
        ("target_start", "Target start"),
        ("target_end", "Target finish"),
        ("act_start", "Actual start"),
        ("act_end", "Actual finish"),
    ]
    _sheet_field_changes(
        wb,
        c,
        "Variances",
        fields,
        "Stored values; duration and float use each activity's calendar hours/day",
    )


def _sheet_logic(wb, c: Comparison) -> None:
    _sheet_relationships(
        wb,
        c,
        "Logic",
        c.logic_added + c.logic_deleted + c.logic_modified,
        include_status=True,
    )


def _sheet_dcma(wb, c: Comparison) -> None:
    ws = wb.create_sheet("DCMA 14")
    _banner(
        ws, "Conventional DCMA 14-point thresholds  ·  stored-value implementation", 8
    )
    _style_header_row(
        ws,
        [
            "#",
            "Metric",
            "Left value",
            "Right value",
            "Threshold",
            "Left",
            "Right",
            "Note",
        ],
        row=4,
    )
    left_m = {m.number: m for m in assess(c.left)}
    right_m = {m.number: m for m in assess(c.right)}
    for n in range(1, 15):
        lm = left_m[n]
        rm = right_m[n]
        row = 4 + n
        ws.cell(row, 1, n)
        ws.cell(row, 2, lm.name)
        ws.cell(row, 3, lm.value)
        ws.cell(row, 4, rm.value)
        ws.cell(row, 5, lm.threshold)
        ws.cell(row, 6, _pass(lm.passed))
        ws.cell(row, 7, _pass(rm.passed))
        ws.cell(row, 8, lm.note or rm.note)
        fill = _dcma_fill(lm.passed, rm.passed)
        _paint_row(ws, row, 8, fill)
        ws.cell(row, 6).fill = _one_fill(lm.passed)
        ws.cell(row, 7).fill = _one_fill(rm.passed)


def _sheet_warnings(wb, c: Comparison) -> None:
    ws = wb.create_sheet("Warnings")
    _banner(ws, "Parser, match, and project-identity warnings", 1)
    _style_header_row(ws, ["Warning"], row=4)
    for i, w in enumerate(c.warnings, start=5):
        ws.cell(i, 1, w)
        _paint_row(ws, i, 1, MOD_FILL)


def _dcma_fill(a: bool | None, b: bool | None) -> PatternFill:
    if a is False or b is False:
        return FAIL_FILL
    if a is True and b is True:
        return PASS_FILL
    return NA_FILL


def _one_fill(v: bool | None) -> PatternFill:
    if v is True:
        return PASS_FILL
    if v is False:
        return FAIL_FILL
    return NA_FILL


def _pass(v: bool | None) -> str:
    if v is True:
        return "PASS"
    if v is False:
        return "FAIL"
    return "N/A"


def _sheet_relationships(wb, c: Comparison, title, items, include_status=False):
    rows = []
    for delta in items:
        schedule = c.left if delta.status == "deleted" else c.right
        activity = schedule.activities.get(delta.succ_code)
        activity = activity or Activity("", delta.succ_code, "Unresolved activity")
        context = [delta.pred_code, delta.rel_type, "revision"]
        if include_status:
            context.append(delta.status)
        rows.append((activity, [(delta.lag_before, delta.lag_after)], context))
    headers = ["Predecessor", "Type", "Class"] + (["Status"] if include_status else [])
    _paired_sheet(
        wb,
        title,
        ["Lag (hours)"],
        rows,
        headers,
        "Activity ID is the successor; lag is stored hours",
    )


def _sheet_driving_path(wb, c: Comparison):
    def membership(activity):
        if activity is None:
            return "Absent"
        if activity.total_float_hr is None:
            return "Unknown"
        return "Member" if activity.is_critical else "Not member"

    rows = []
    for code in sorted(set(c.left.activities) | set(c.right.activities)):
        left, right = c.left.activities.get(code), c.right.activities.get(code)
        before, after = membership(left), membership(right)
        if "Member" not in {before, after}:
            continue
        if left is None:
            change = "Added task"
        elif right is None:
            change = "Deleted task"
        elif "Unknown" in {before, after}:
            change = "Stored float availability changed"
        elif before == after:
            change = "Retained"
        else:
            change = "Entered" if after == "Member" else "Left"
        rows.append(
            (
                right or left,
                [
                    (before, after),
                    (
                        left.total_float_days if left else None,
                        right.total_float_days if right else None,
                    ),
                ],
                [change],
            )
        )
    _paired_sheet(
        wb,
        "Driving Path Change",
        ["Membership", "Total float (days)"],
        rows,
        ["Change type"],
        "Approximation from stored TF <= 0 hours; not a calculated driving path",
    )


def _sheet_resources(wb, c: Comparison, title, assignments, added):
    fields = [
        ("rsrc_name", "Resource"),
        ("target_qty", "Target quantity"),
        ("remain_qty", "Remaining quantity"),
        ("act_reg_qty", "Actual regular quantity"),
        ("target_cost", "Target cost"),
        ("remain_cost", "Remaining cost"),
        ("act_reg_cost", "Actual regular cost"),
    ]
    schedule = c.right if added else c.left
    rows = []
    for assignment in sorted(
        assignments, key=lambda a: (a.task_code, a.rsrc_name, a.rsrc_id)
    ):
        activity = schedule.activities[assignment.task_code]
        values = [getattr(assignment, field) for field, _ in fields]
        pairs = [(None, value) if added else (value, None) for value in values]
        rows.append((activity, pairs, [assignment.rsrc_id, "revision"]))
    _paired_sheet(
        wb,
        title,
        [label for _, label in fields],
        rows,
        ["Resource ID (context)", "Class"],
        "Assignment inventory; quantities and costs are stored values",
    )


def _sheet_udfs(wb, c: Comparison, title, deltas):
    rows = []
    for delta in deltas:
        activity = (
            c.right.activities.get(delta.task_code)
            or c.left.activities[delta.task_code]
        )
        names = (
            delta.name if delta.status != "added" else None,
            delta.name if delta.status != "deleted" else None,
        )
        rows.append((activity, [names, (delta.before, delta.after)], ["revision"]))
    ws = _paired_sheet(
        wb,
        title,
        ["UDF", "Value"],
        rows,
        ["Class"],
        "Activity UDFs matched by label (or name when no label is exported)",
    )
    _format_udf_rows(
        ws,
        c,
        ((row, delta.task_code, delta.name) for row, delta in enumerate(deltas, 5)),
        base_column=6,
    )


def _sheet_combined(wb, c: Comparison):
    rows = []

    def add(code, before, after, change_type, field, klass="revision", deleted=False):
        primary, secondary = (c.left, c.right) if deleted else (c.right, c.left)
        activity = primary.activities.get(code) or secondary.activities.get(code)
        activity = activity or Activity("", code, "Unresolved activity")
        rows.append((activity, [(before, after)], [change_type, field, klass]))

    for delta in c.added:
        add(delta.task_code, None, delta.task_code, "Added Task", "Activity ID")
    for delta in c.deleted:
        add(
            delta.task_code,
            delta.task_code,
            None,
            "Deleted Task",
            "Activity ID",
            deleted=True,
        )
    for delta in c.modified:
        for change in delta.changes:
            if change.field.startswith("udf:"):
                continue  # The UDF inventory below represents this event once.
            add(
                delta.task_code,
                change.before,
                change.after,
                "Activity Field",
                change.field,
                delta.klass,
            )
    for deltas, label in [
        (c.logic_added, "Added Relationship"),
        (c.logic_deleted, "Deleted Relationship"),
        (c.logic_modified, "Lag Change"),
    ]:
        for delta in deltas:
            add(
                delta.succ_code,
                delta.lag_before,
                delta.lag_after,
                label,
                f"{delta.pred_code} -> {delta.succ_code} ({delta.rel_type}); lag hours",
                deleted=delta.status == "deleted",
            )
    for assignments, added in [(c.resource_added, True), (c.resource_deleted, False)]:
        for assignment in assignments:
            name = assignment.rsrc_name or assignment.rsrc_id
            add(
                assignment.task_code,
                None if added else name,
                name if added else None,
                "Added Resource Assignment" if added else "Deleted Resource Assignment",
                "Resource",
                deleted=not added,
            )
    for deltas, label in [
        (c.udf_added, "Added UDF"),
        (c.udf_deleted, "Deleted UDF"),
        (c.udf_modified, "Revised UDF"),
    ]:
        for delta in deltas:
            add(
                delta.task_code,
                delta.before,
                delta.after,
                label,
                delta.name,
                deleted=delta.status == "deleted",
            )
    for delta in c.resource_modified:
        for change in delta.changes:
            add(
                delta.after.task_code,
                change.before,
                change.after,
                "Resource Value Change",
                f"{delta.after.rsrc_name}: {RESOURCE_FIELDS[change.field]}",
            )
    for delta in c.calendar_changes:
        add(
            delta.task_code,
            delta.before,
            delta.after,
            "Calendar Definition",
            delta.field,
        )
    rows.sort(key=lambda row: (row[0].task_code, row[2][0], row[2][1]))
    ws = _paired_sheet(
        wb,
        "Changes Combined",
        ["Value"],
        rows,
        ["Change type", "Field", "Class"],
        "One row per stored-field or inventory change; derived views are not repeated",
    )
    for row, (_, _, context) in enumerate(rows, 5):
        if context[0] == "Resource Value Change":
            for col in (4, 5):
                ws.cell(row, col).number_format = (
                    COST_FORMAT
                    if "cost" in context[1].split(": ")[-1].lower()
                    else NUMBER_FORMAT
                )
    _format_udf_rows(
        ws,
        c,
        (
            (row, activity.task_code, context[1])
            for row, (activity, _, context) in enumerate(rows, 5)
            if context[0] in {"Added UDF", "Deleted UDF", "Revised UDF"}
        ),
    )


def _summary_catalog(wb):
    ws = wb["Summary"]
    ws.merge_cells("A33:C33")
    ws["A33"] = "REPORT CATALOG · counts are report rows"
    ws["A33"].fill, ws["A33"].font = NAVY_FILL, HEADER_FONT
    for col, label in enumerate(["Group", "Report (click to open)", "Rows"], 1):
        ws.cell(34, col, label).fill = STEEL_FILL
        ws.cell(34, col).font = HEADER_FONT
    for row, detail in enumerate(
        (s for s in wb if s.title not in {"Summary", "Legend"}), 35
    ):
        title = detail.title
        if title in {
            "DCMA 14",
            "Warnings",
            "Quality Checks",
            "Quality Findings",
            "Open Ends",
            "Constraint Register",
            "Out of Sequence",
            "Date Review",
            "Check Settings",
        }:
            group = "Quality"
        elif "Relationship" in title or title in {
            "Lag Change",
            "Logic",
            "Driving Path Change",
        }:
            group = "Network"
        elif title in {"Actual Date Change", "Progress Change"}:
            group = "Progress"
        elif title in {
            "Variances",
            "Float Change",
            "Milestone Slips",
        } or title.startswith(("Early", "Late")):
            group = "Dates and float"
        elif title in {"Modified", "Changes Combined"}:
            group = "Combined detail"
        else:
            group = "Revisions and inventory"
        ws.cell(row, 1, group)
        link = ws.cell(row, 2, title)
        link.hyperlink = Hyperlink(
            ref=link.coordinate, location=f"'{title}'!A1", display=title
        )
        ws.cell(row, 3, max(0, detail.max_row - 4))
        _paint_row(ws, row, 3, ZEBRA if row % 2 else WHITE)
        link.font = Font(name="Calibri", size=11, color=STEEL, underline="single")
    ws.freeze_panes = "A5"


def _sheet_advanced_values(wb, c):
    rows = [
        (c.right.activities[d.task_code], [(d.before, d.after)], [d.subject, d.field])
        for d in c.calendar_changes
    ]
    _paired_sheet(
        wb,
        "Calendar Definitions",
        ["Value"],
        rows,
        ["Calendar", "Field"],
        "Effective inherited workweek and exceptions for shared activities; no CPM",
    )
    rows = []
    for d in c.resource_modified:
        for change in d.changes:
            rows.append(
                (
                    c.right.activities[d.after.task_code],
                    [(change.before, change.after)],
                    [d.after.rsrc_name, RESOURCE_FIELDS[change.field]],
                )
            )
    ws = _paired_sheet(
        wb,
        "Resource Value Changes",
        ["Value"],
        rows,
        ["Resource", "Field"],
        "Stored assignment quantities, costs and units; see warnings for ambiguous matches",
    )
    for row, (_, _, context) in enumerate(rows, 5):
        for col in (4, 5):
            ws.cell(row, col).number_format = (
                COST_FORMAT if "cost" in context[1].lower() else NUMBER_FORMAT
            )
    # Changes Combined is built after this sheet and formats these values there.
