"""Local update history with an explicit order and optional baseline."""

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4
from collections import Counter

from xercompare.model.schedule import build_schedule
from xercompare.parse.xer_parser import parse_xer
from xercompare.compare.engine import compare_schedules
from xercompare.quality.checks import parsed_date, inspect_schedule
from xercompare.quality.dcma import assess
from xercompare.quality.profile import QualityProfile


@dataclass(frozen=True)
class Snapshot:
    path: Path
    project: str
    data_date: str
    activity_count: int


def discover_history(folder):
    folder = Path(folder)
    if not folder.is_dir():
        raise ValueError("Choose an existing folder of XER exports.")
    snapshots = []
    for path in sorted(folder.iterdir()):
        if path.is_file() and path.suffix.lower() == ".xer":
            s = build_schedule(parse_xer(path))
            if parsed_date(s.project.data_date) is None:
                raise ValueError(
                    f"{path.name}: a valid stored data date is required for history."
                )
            snapshots.append(
                Snapshot(
                    path.resolve(),
                    s.project.short_name,
                    s.project.data_date,
                    s.activity_count,
                )
            )
    return sorted(
        snapshots, key=lambda s: (parsed_date(s.data_date), s.path.name.casefold())
    )


@dataclass(frozen=True)
class HistoryRequest:
    paths: tuple
    output: Path
    baseline: Path | None = None
    profile: QualityProfile | None = None
    allow_project_names: bool = False

    def validate(self):
        if len(self.paths) < 2 or len({Path(p).resolve() for p in self.paths}) != len(
            self.paths
        ):
            raise ValueError("Choose at least two distinct XER exports for history.")
        for path in [*self.paths, *([self.baseline] if self.baseline else [])]:
            if not Path(path).is_file() or Path(path).suffix.lower() != ".xer":
                raise ValueError(f"Choose an existing XER: {path}")
        if self.output.exists() and not self.output.is_dir():
            raise ValueError("Choose an output folder, not a file.")


@dataclass
class HistoryPair:
    kind: str
    base: str
    revised: str
    summary: dict
    report: str


@dataclass
class HistoryResult:
    folder: Path
    excel: Path
    pdf: Path
    trends: list
    pairs: list
    warnings: list


def run_history(request, progress=None, cancel=None):
    from xercompare.gui.service import Cancelled
    from xercompare.reports.excel import write_excel
    from xercompare.reports.pdf import write_pdf

    request.validate()

    def step(value, message):
        if progress:
            progress(value, message)
        if cancel is not None and cancel.is_set():
            raise Cancelled()

    schedules = []
    for i, path in enumerate(request.paths):
        step(
            int(20 * i / len(request.paths)),
            f"Reading update {i+1} of {len(request.paths)}…",
        )
        s = build_schedule(parse_xer(path))
        if parsed_date(s.project.data_date) is None:
            raise ValueError(
                f"{Path(path).name}: a valid stored data date is required."
            )
        schedules.append(s)
    baseline = build_schedule(parse_xer(request.baseline)) if request.baseline else None
    names = {
        s.project.short_name for s in [*schedules, *([baseline] if baseline else [])]
    }
    if len(names) > 1 and not request.allow_project_names:
        raise ValueError(
            "Different project names: confirm that these exports belong to the same project before running history."
        )
    warnings = [f"{s.source.name}: {w}" for s in schedules for w in s.warnings]
    dates = [s.project.data_date for s in schedules]
    if len(set(dates)) < len(dates):
        warnings.append(
            "Updates share the same data date. Reviewed order is used; equal dates do not establish chronology."
        )
    if dates != sorted(dates):
        warnings.append("The reviewed order is not chronological by data date.")
    if len(names) > 1:
        warnings.append(
            "Different project names were explicitly accepted for this history."
        )
    profile = request.profile or QualityProfile()
    trends = []
    for s in schedules:
        finishes = [
            parsed_date(a.early_end)
            for a in s.activities.values()
            if parsed_date(a.early_end)
        ]
        metrics = assess(s)
        checks = inspect_schedule(s, profile)
        trends.append(
            [
                s.source.name,
                s.project.short_name,
                s.project.data_date,
                s.activity_count,
                max(finishes).strftime("%Y-%m-%d %H:%M") if finishes else "",
                sum(a.is_critical for a in s.activities.values()),
                sum(m.passed is False for m in metrics),
                sum(m.passed is None for m in metrics),
                sum(r.passed is False for r in checks),
                sum(r.passed is None for r in checks),
            ]
        )
    jobs = [("Consecutive", a, b) for a, b in zip(schedules, schedules[1:])]
    if baseline:
        jobs.extend(
            ("Baseline", baseline, s)
            for s in schedules
            if s.source.resolve() != baseline.source.resolve()
        )
    request.output.mkdir(parents=True, exist_ok=True)
    folder = (
        request.output / f"XERHistory-{datetime.now():%Y%m%d-%H%M%S}-{uuid4().hex[:6]}"
    )
    with TemporaryDirectory(prefix=".xerhistory-", dir=request.output) as temporary:
        stage = Path(temporary)
        pairs = []
        for i, (kind, left, right) in enumerate(jobs):
            step(
                25 + int(65 * i / len(jobs)),
                f"Comparing {i+1} of {len(jobs)}: {left.source.name} → {right.source.name}",
            )
            c = compare_schedules(left, right, profile)
            relative = f"{kind}/{i+1:03d}"
            write_excel(c, stage / relative / "compare.xlsx")
            step(25 + int(65 * (i + 0.5) / len(jobs)), "Writing comparison summary…")
            write_pdf(c, stage / relative / "compare.pdf")
            pairs.append(
                HistoryPair(
                    kind, left.source.name, right.source.name, c.summary, relative
                )
            )
            warnings.extend(
                f"{kind} {i+1}: {w}"
                for w in c.warnings
                if w not in left.warnings + right.warnings
            )
        step(92, "Writing history trends and charts…")
        _write_history(stage, trends, pairs, warnings, profile, request.baseline)
        step(99, "Finishing history reports…")
        stage.rename(folder)
    if progress:
        progress(100, "History complete. Trends and comparison reports are ready.")
    return HistoryResult(
        folder, folder / "history.xlsx", folder / "history.pdf", trends, pairs, warnings
    )


def _write_history(folder, trends, pairs, warnings, profile, baseline):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from xercompare.reports.excel import (
        _banner,
        _style_header_row,
        _paint_row,
        ZEBRA,
        WHITE,
    )
    from xercompare.reports.formatting import NUMBER_FORMAT
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib import colors
    from html import escape

    wb = Workbook()
    wb.remove(wb.active)

    def grid(title, headers, rows):
        ws = wb.create_sheet(title)
        _banner(ws, title, len(headers))
        _style_header_row(ws, headers)
        for i, row in enumerate(rows, 5):
            for col, value in enumerate(row, 1):
                cell = ws.cell(i, col, value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if isinstance(value, float):
                    cell.number_format = NUMBER_FORMAT
            _paint_row(ws, i, len(headers), ZEBRA if i % 2 else WHITE)
        ws.auto_filter.ref = f"A4:{ws.cell(max(ws.max_row,4),len(headers)).coordinate}"
        return ws

    grid(
        "History Summary",
        ["Setting", "Value"],
        [
            ["Profile", profile.name],
            ["Baseline", str(baseline) if baseline else "None selected"],
            ["Order", "User-reviewed order; defaults to data date then filename"],
            ["Finish", "Latest stored activity early finish; not recalculated"],
            ["Critical count", "Stored total float <= 0 membership approximation"],
            [
                "Quality",
                "Flag counts depend on data availability; N/A counts shown beside them",
            ],
            [
                "Details",
                "Open a linked workbook on Comparisons for the full changes and PDF beside it.",
            ],
            ["High float days", profile.high_float_days],
            ["High duration days", profile.high_duration_days],
        ],
    )
    headers = [
        "File",
        "Project",
        "Data date",
        "Activities",
        "Latest stored early finish",
        "Stored TF <= 0 count",
        "DCMA flags",
        "DCMA N/A",
        "Profile flags",
        "Profile not assessed",
    ]
    ws = grid("Trends", headers, trends)
    ws.column_dimensions["A"].width = 54
    for row, data in enumerate(trends, 5):
        if data[4]:
            ws.cell(row, 5, parsed_date(data[4])).number_format = "yyyy-mm-dd hh:mm"
    for index, column in enumerate([4, 5, 6, 9]):
        chart = LineChart()
        chart.title = headers[column - 1]
        chart.y_axis.title = headers[column - 1]
        chart.x_axis.title = "Updates in reviewed order"
        if column == 5:
            chart.y_axis.numFmt = "yyyy-mm-dd"
        chart.add_data(
            Reference(ws, min_col=column, min_row=4, max_row=ws.max_row),
            titles_from_data=True,
        )
        chart.set_categories(Reference(ws, min_col=1, min_row=5, max_row=ws.max_row))
        chart.width = 24
        chart.height = 9
        ws.add_chart(chart, f"L{5+index*19}")
    rows = [
        [
            p.kind,
            p.base,
            p.revised,
            p.summary["added"],
            p.summary["deleted"],
            p.summary["modified"],
            p.summary["logic_added"]
            + p.summary["logic_deleted"]
            + p.summary["logic_modified"],
            p.report + "/compare.xlsx",
        ]
        for p in pairs
    ]
    ws = grid(
        "Comparisons",
        [
            "Comparison",
            "Base file",
            "Revised file",
            "Added",
            "Deleted",
            "Modified",
            "Relationship edits",
            "Open full workbook",
        ],
        rows,
    )
    for row, p in enumerate(pairs, 5):
        ws.cell(row, 8).hyperlink = p.report + "/compare.xlsx"
    grid("Warnings", ["Warning"], [[w] for w in warnings])
    wb.save(folder / "history.xlsx")
    styles = getSampleStyleSheet()
    story = [
        Paragraph("XERCOMPARE · Update history", styles["Title"]),
        Paragraph(
            f"{len(trends)} exports; {len(pairs)} complete comparison sets. Profile: {escape(profile.name)}.",
            styles["BodyText"],
        ),
        Paragraph(
            "Order is the reviewed sequence. Latest finish is stored early finish; critical count is stored TF <= 0. No CPM recalculation.",
            styles["BodyText"],
        ),
        Spacer(1, 12),
    ]
    rows = [
        [
            "Update",
            "Data date",
            "Activities",
            "Latest finish",
            "TF <= 0",
            "DCMA flags / N/A",
        ]
    ]
    for r in trends:
        rows.append(
            [
                Paragraph(escape(r[0]), styles["BodyText"]),
                r[2][:10],
                r[3],
                r[4][:16],
                r[5],
                f"{r[6]} / {r[7]}",
            ]
        )
    table = Table(rows, colWidths=[230, 80, 65, 110, 65, 140], repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D6E3F0")),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )
    story.append(table)
    story.append(Spacer(1, 12))
    for warning in warnings[:12]:
        story.append(Paragraph(escape(warning), styles["BodyText"]))
    story.append(
        Paragraph(
            "Open history.xlsx for charts and links to each full comparison workbook. Each comparison folder also contains its PDF summary.",
            styles["BodyText"],
        )
    )
    SimpleDocTemplate(
        str(folder / "history.pdf"),
        pagesize=landscape(letter),
        leftMargin=36,
        rightMargin=36,
    ).build(story)
