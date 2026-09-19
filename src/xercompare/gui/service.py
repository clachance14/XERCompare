"""Desktop orchestration around the existing parser, comparison, and reports."""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from threading import Event
from typing import Callable
from uuid import uuid4

from openpyxl import load_workbook

from xercompare.compare.engine import compare_schedules
from xercompare.model.schedule import build_schedule
from xercompare.parse.xer_parser import parse_xer
from xercompare.quality.dcma import Metric, assess
from xercompare.quality.checks import inspect_schedule
from xercompare.quality.profile import QualityProfile
from xercompare.reports.excel import write_excel
from xercompare.reports.pdf import write_pdf
from xercompare.reports.formatting import display_value


class Cancelled(Exception):
    """Cancellation between stages; no incomplete reports are published."""


@dataclass(frozen=True)
class CompareRequest:
    left: Path
    right: Path
    output: Path
    profile: QualityProfile | None = None

    def validate(self):
        for label, path in [("Older XER", self.left), ("Revised XER", self.right)]:
            if not path.is_file():
                raise ValueError(f"{label}: choose an existing .xer file.")
            if path.suffix.lower() != ".xer":
                raise ValueError(f"{label}: choose a Primavera P6 .xer file.")
        if self.output.exists() and not self.output.is_dir():
            raise ValueError("Output folder: choose a folder, not a file.")


@dataclass(frozen=True)
class ReportInfo:
    name: str
    rows: int | None


@dataclass
class RunResult:
    folder: Path
    excel: Path
    pdf: Path
    summary: dict[str, int]
    reports: list[ReportInfo]
    quality: list[Metric]
    warnings: list[str]
    left_project: str
    right_project: str
    left_date: str
    right_date: str
    resource_added: int
    resource_deleted: int
    udf_added: int
    udf_deleted: int
    udf_modified: int
    resource_modified: int = 0
    checks: list = field(default_factory=list)
    base_checks: list = field(default_factory=list)


@dataclass
class Preview:
    headers: list[str]
    rows: list[tuple]
    total: int


def read_preview(
    path: Path, sheet: str, limit: int = 200, filters=None, query=""
) -> Preview:
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = book[sheet]
        if sheet == "Summary":
            headers = ["Field", "Value", "Field 2", "Value 2", "Field 3", "Value 3"]
            first = 4
        else:
            headers = [
                str(cell.value or "")
                for cell in next(ws.iter_rows(min_row=4, max_row=4))
            ]
            first = 5
        if filters or query:
            rows = []
            total = 0
            for row in ws.iter_rows(min_row=first):
                values = tuple(
                    display_value(cell.value, cell.number_format) for cell in row
                )
                mapping = dict(zip(headers, values))
                if any(mapping.get(k) != v for k, v in (filters or {}).items()):
                    continue
                if query and not any(
                    query.casefold() in str(v).casefold() for v in values
                ):
                    continue
                total += 1
                if len(rows) < limit:
                    rows.append(values)
            return Preview(headers, rows, total)
        total = max(0, (ws.max_row or 0) - first + 1)
        rows = (
            [
                tuple(display_value(cell.value, cell.number_format) for cell in row)
                for row in ws.iter_rows(
                    min_row=first,
                    max_row=min(ws.max_row, first + limit - 1),
                )
            ]
            if total
            else []
        )
        return Preview(headers, rows, total)
    finally:
        book.close()


def run_comparison(
    request: CompareRequest,
    progress: Callable[[int, str], None] | None = None,
    cancel: Event | None = None,
) -> RunResult:
    request.validate()

    def step(value, message):
        if progress:
            progress(value, message)
        if cancel is not None and cancel.is_set():
            raise Cancelled()

    step(5, "Reading the older XER...")
    left = build_schedule(parse_xer(request.left))
    step(20, "Reading the revised XER...")
    right = build_schedule(parse_xer(request.right))
    step(40, "Comparing activities, relationships, resources, and UDFs...")
    comparison = compare_schedules(left, right, request.profile)
    request.output.mkdir(parents=True, exist_ok=True)
    folder = (
        request.output / f"XERCompare-{datetime.now():%Y%m%d-%H%M%S}-{uuid4().hex[:6]}"
    )
    with TemporaryDirectory(prefix=".xercompare-", dir=request.output) as temporary:
        stage = Path(temporary)
        step(55, "Writing the Excel report catalog...")
        write_excel(comparison, stage / "compare.xlsx")
        step(80, "Writing the PDF summary...")
        write_pdf(comparison, stage / "compare.pdf")
        step(95, "Preparing the report previews...")
        book = load_workbook(stage / "compare.xlsx", read_only=True, data_only=True)
        try:
            reports = [
                ReportInfo(
                    ws.title, None if ws.title == "Summary" else max(0, ws.max_row - 4)
                )
                for ws in book
            ]
        finally:
            book.close()
        quality = assess(right)
        # Both reports are complete before the dated output folder becomes visible.
        step(98, "Finishing the report set...")
        stage.rename(folder)
    if progress:
        progress(100, "Comparison complete. Your reports are ready.")
    return RunResult(
        folder,
        folder / "compare.xlsx",
        folder / "compare.pdf",
        comparison.summary,
        reports,
        quality,
        comparison.warnings,
        left.project.short_name,
        right.project.short_name,
        left.project.data_date,
        right.project.data_date,
        len(comparison.resource_added),
        len(comparison.resource_deleted),
        len(comparison.udf_added),
        len(comparison.udf_deleted),
        len(comparison.udf_modified),
        len(comparison.resource_modified),
        inspect_schedule(right, request.profile),
        inspect_schedule(left, request.profile),
    )
